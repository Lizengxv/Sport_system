import io



import os
import hashlib
import re



import random
import math
from copy import copy



from datetime import datetime



from typing import List, Optional, Dict







import pandas as pd
import urllib.parse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Side
from openpyxl.utils import get_column_letter



from fastapi import FastAPI, Depends, HTTPException, UploadFile, File



from fastapi.middleware.cors import CORSMiddleware



from fastapi.responses import StreamingResponse



from passlib.context import CryptContext



from sqlalchemy.orm import Session



from sqlalchemy import select, delete, distinct, text







from .db import Base, engine, SessionLocal



from .models import (



    User,



    Athlete,



    Grouping,



    Event,



    Result,



    ResultDistance,



    ResultJump,



    PersonalRanking,



    CollegeRanking,
    TeamRankingManualBonus,



)



from .schemas import (



    LoginRequest,



    RegisterRequest,



    ResetPasswordRequest,



    AthleteCreate,



    AthleteUpdate,



    GroupGenerateRequest,



    GroupConfirmRequest,



    ResultSubmitRequest,



    ResultDistanceSubmitRequest,
    TeamRankingManualBonusRequest,



)











Base.metadata.create_all(bind=engine)


def ensure_result_unique_indexes():
    # Keep old deployments compatible: legacy schema may still have uq_result_event(student_id,event).
    with engine.begin() as conn:
        try:
            conn.execute(text("ALTER TABLE results DROP INDEX uq_result_event"))
        except Exception:
            pass
        try:
            conn.execute(
                text(
                    "ALTER TABLE results "
                    "ADD UNIQUE INDEX uq_result_event_round (student_id, event, round)"
                )
            )
        except Exception:
            pass


ensure_result_unique_indexes()







app = FastAPI(title="School Sports Management System")







app.add_middleware(



    CORSMiddleware,



    allow_origins=["*"],



    allow_credentials=True,



    allow_methods=["*"],



    allow_headers=["*"],



)







pwd_context = CryptContext(schemes=["bcrypt"], deprecated="auto")

LAST_IMPORT_TS = None







MULTI_ROUND_EVENTS = {
    "\u0031\u0030\u0030\u7c73",
    "\u0032\u0030\u0030\u7c73",
    "\u0034\u0030\u0030\u7c73",
    "\u0038\u0030\u0030\u7c73",
    "\u0031\u0035\u0030\u0030\u7c73",
    "\u0033\u0030\u0030\u0030\u7c73",
    "\u0035\u0030\u0030\u0030\u7c73",
    "\u0034*\u0031\u0030\u0030\u7c73",
    "\u0034*\u0032\u0030\u0030\u7c73",
    "\u0034*\u0034\u0030\u0030\u7c73",
}
MIXED_EVENTS = {"\u6df7\u5408\u0034*\u0031\u0030\u0030\u7c73", "\u6df7\u5408\u0034*\u0034\u0030\u0030\u7c73"}
DISTANCE_KEYWORDS = {
    "\u8df3",
    "\u6295",
    "\u63b7",
    "\u94c5\u7403",
    "\u6807\u67aa",
    "\u94c1\u997c",
    "\u8df3\u8fdc",
    "\u4e09\u7ea7\u8df3\u8fdc",
    "\u7acb\u5b9a\u8df3\u8fdc",
    "\u6025\u884c\u8df3\u8fdc",
    "\u8df3\u9ad8",
}
DISTANCE_EVENTS_ALIAS = {
    "\u8df3\u8fdc",
    "\u4e09\u7ea7\u8df3\u8fdc",
    "\u7acb\u5b9a\u8df3\u8fdc",
    "\u6025\u884c\u8df3\u8fdc",
    "\u94c5\u7403",
    "\u6807\u67aa",
    "\u94c1\u997c",
}














def get_db():



    db = SessionLocal()



    try:



        yield db



    finally:



        db.close()











def verify_password(plain, hashed):



    return pwd_context.verify(plain, hashed)











def hash_password(plain):



    return pwd_context.hash(plain)



















DISTANCE_EVENTS_ALIAS = {"璺宠繙", "绔嬪畾璺宠繙", "鎬ヨ璺宠繙", "閾呯悆"}











# Override garbled defaults with normalized UTF-8 values.

MULTI_ROUND_EVENTS = {
    "\u0031\u0030\u0030\u7c73",
    "\u0032\u0030\u0030\u7c73",
    "\u0034\u0030\u0030\u7c73",
    "\u0038\u0030\u0030\u7c73",
    "\u0031\u0035\u0030\u0030\u7c73",
    "\u0033\u0030\u0030\u0030\u7c73",
    "\u0035\u0030\u0030\u0030\u7c73",
    "\u0034*\u0031\u0030\u0030\u7c73",
    "\u0034*\u0032\u0030\u0030\u7c73",
    "\u0034*\u0034\u0030\u0030\u7c73",
}
MIXED_EVENTS = {"\u6df7\u5408\u0034*\u0031\u0030\u0030\u7c73", "\u6df7\u5408\u0034*\u0034\u0030\u0030\u7c73"}
DISTANCE_KEYWORDS = {
    "\u8df3",
    "\u6295",
    "\u63b7",
    "\u94c5\u7403",
    "\u6807\u67aa",
    "\u94c1\u997c",
    "\u8df3\u8fdc",
    "\u4e09\u7ea7\u8df3\u8fdc",
    "\u7acb\u5b9a\u8df3\u8fdc",
    "\u6025\u884c\u8df3\u8fdc",
    "\u8df3\u9ad8",
}
DISTANCE_EVENTS_ALIAS = {
    "\u8df3\u8fdc",
    "\u4e09\u7ea7\u8df3\u8fdc",
    "\u7acb\u5b9a\u8df3\u8fdc",
    "\u6025\u884c\u8df3\u8fdc",
    "\u94c5\u7403",
    "\u6807\u67aa",
    "\u94c1\u997c",
}


# Override garbled defaults with normalized UTF-8 values.



MULTI_ROUND_EVENTS = {
    "\u0031\u0030\u0030\u7c73",
    "\u0032\u0030\u0030\u7c73",
    "\u0034\u0030\u0030\u7c73",
    "\u0038\u0030\u0030\u7c73",
    "\u0031\u0035\u0030\u0030\u7c73",
    "\u0033\u0030\u0030\u0030\u7c73",
    "\u0035\u0030\u0030\u0030\u7c73",
    "\u0034*\u0031\u0030\u0030\u7c73",
    "\u0034*\u0032\u0030\u0030\u7c73",
    "\u0034*\u0034\u0030\u0030\u7c73",
}
MIXED_EVENTS = {"\u6df7\u5408\u0034*\u0031\u0030\u0030\u7c73", "\u6df7\u5408\u0034*\u0034\u0030\u0030\u7c73"}
DISTANCE_KEYWORDS = {
    "\u8df3",
    "\u6295",
    "\u63b7",
    "\u94c5\u7403",
    "\u6807\u67aa",
    "\u94c1\u997c",
    "\u8df3\u8fdc",
    "\u4e09\u7ea7\u8df3\u8fdc",
    "\u7acb\u5b9a\u8df3\u8fdc",
    "\u6025\u884c\u8df3\u8fdc",
    "\u8df3\u9ad8",
}
DISTANCE_EVENTS_ALIAS = {
    "\u8df3\u8fdc",
    "\u4e09\u7ea7\u8df3\u8fdc",
    "\u7acb\u5b9a\u8df3\u8fdc",
    "\u6025\u884c\u8df3\u8fdc",
    "\u94c5\u7403",
    "\u6807\u67aa",
    "\u94c1\u997c",
}










def best_distance(*values):



    vals = [v for v in values if v is not None]



    return max(vals) if vals else None











def normalize_event_name(event_name: str) -> str:
    value = (event_name or "").strip().lower()
    value = value.replace(" ", "")
    value = value.replace("x", "*").replace("\u00d7", "*").replace("\uff0a", "*")
    return value


def normalize_gender_value(gender: Optional[str]) -> str:
    value = (gender or "").strip()
    lowered = value.lower()
    if value in {"男", "男子", "男子组"} or lowered in {"male", "m"}:
        return "男"
    if value in {"女", "女子", "女子组"} or lowered in {"female", "f"}:
        return "女"
    if value in {"混合", "混合组"} or lowered in {"mixed", "mix"}:
        return "混合"
    return value


def gender_filter_values(gender: Optional[str]) -> List[str]:
    normalized = normalize_gender_value(gender)
    if not normalized:
        return []
    if normalized == "男":
        return ["男", "男子", "男子组"]
    if normalized == "女":
        return ["女", "女子", "女子组"]
    if normalized == "混合":
        return ["混合", "混合组"]
    return [normalized]


def is_high_jump_event(event_name: str) -> bool:
    value = normalize_event_name(event_name)
    keywords = ["\u8df3\u9ad8", "highjump", "high-jump"]
    return any(k in value for k in keywords)


def is_distance_field_event(event_name: str) -> bool:
    if is_high_jump_event(event_name):
        return False
    value = normalize_event_name(event_name)
    keywords = [
        "\u8df3\u8fdc",
        "\u4e09\u7ea7\u8df3\u8fdc",
        "\u7acb\u5b9a\u8df3\u8fdc",
        "\u6025\u884c\u8df3\u8fdc",
        "\u94c5\u7403",
        "\u6807\u67aa",
        "\u94c1\u997c",
        "\u94fe\u7403",
        "longjump",
        "triplejump",
        "standinglongjump",
        "shotput",
        "javelin",
        "discus",
        "hammerthrow",
    ]
    return any(k in value for k in keywords)


def classify_event_table(event_name: str) -> str:
    if is_high_jump_event(event_name):
        return "jump"
    if is_distance_field_event(event_name):
        return "distance"
    return "track"


def is_relay_event(event_name: str) -> bool:
    value = normalize_event_name(event_name)
    relay_keywords = ["4*100", "4*200", "4*400", "\u63a5\u529b", "relay"]
    return any(keyword in value for keyword in relay_keywords)


def is_long_distance_track_event(event_name: Optional[str]) -> bool:
    if not event_name or classify_event_table(event_name) != "track" or is_relay_event(event_name):
        return False
    value = normalize_event_name(event_name)
    match = re.search(r"(\d+)", value)
    if not match:
        return False
    meters = int(match.group(1))
    return meters > 800


def build_relay_storage_student_id(event_name: str, name: str, gender: str = "") -> str:
    raw_key = f"{normalize_event_name(event_name)}|{(name or '').strip()}|{(gender or '').strip()}"
    digest = hashlib.sha1(raw_key.encode("utf-8")).hexdigest()[:20].upper()
    return f"TEAM_{digest}"


def score_higher_better(event_name: str) -> bool:
    kind = classify_event_table(event_name)
    return kind in {"distance", "jump"}


def has_distance_records(event_name: str, db: Session) -> bool:
    row = db.scalar(
        select(ResultDistance.id).where(ResultDistance.event == event_name).limit(1)
    )
    return row is not None


def has_jump_records(event_name: str, db: Session) -> bool:
    row = db.scalar(
        select(ResultJump.id).where(ResultJump.event == event_name).limit(1)
    )
    return row is not None


def is_distance_table_event(event_name: str) -> bool:
    return classify_event_table(event_name) == "distance"


def rank_distance_event(event_name: str, round_value: str, db: Session):



    rows = db.scalars(select(ResultDistance).where(ResultDistance.event == event_name)).all()



    scored = []



    for row in rows:



        best = row.final_best if round_value in {"final", "one"} else row.prelim_best



        if not is_rankable_score(best):



            continue



        scored.append({"id": row.id, "score": best})



    ranked = rank_athletes_by_score(scored, event_name)
    ranks = {item["id"]: item["rank"] for item in ranked}



    for row in rows:



        row.rank = ranks.get(row.id)











def list_distance_results(event_name: str, round_value: str | None, college: str | None, name: str | None, student_id: str | None, gender: str | None, db: Session):
    query = select(ResultDistance).where(ResultDistance.event == event_name)
    if college:
        query = query.where(ResultDistance.college == college)
    if name:
        query = query.where(ResultDistance.name == name)
    if student_id:
        query = query.where(ResultDistance.student_id == student_id)
    if gender:
        athlete_query = select(Athlete.student_id).where(
            Athlete.event == event_name,
            Athlete.gender.in_(gender_filter_values(gender)),
        )
        student_ids = [row[0] for row in db.execute(athlete_query).all()]
        if not student_ids:
            return []
        query = query.where(ResultDistance.student_id.in_(student_ids))
    rows = db.scalars(query).all()

    athlete_map = {}
    if rows:
        athlete_query = select(Athlete).where(
            Athlete.student_id.in_([row.student_id for row in rows]),
            Athlete.event == event_name,
        )
        athlete_map = {
            (row.student_id, row.event): row.record_broken
            for row in db.scalars(athlete_query).all()
        }

    payload = []
    for row in rows:
        base = {
            "id": row.id,
            "event": row.event,
            "name": row.name,
            "student_id": row.student_id,
            "college": row.college,
            "prelim_attempt1": row.prelim_attempt1,
            "prelim_attempt2": row.prelim_attempt2,
            "prelim_attempt3": row.prelim_attempt3,
            "prelim_best": row.prelim_best,
            "final_attempt1": row.final_attempt1,
            "final_attempt2": row.final_attempt2,
            "final_attempt3": row.final_attempt3,
            "final_best": row.final_best,
            "rank": row.rank,
            "record_broken": athlete_map.get((row.student_id, row.event), False),
        }

        if round_value == "prelim":
            if row.prelim_best is not None:
                payload.append({
                    **base,
                    "score": row.prelim_best,
                    "rank": row.rank if is_rankable_score(row.prelim_best) else None,
                    "date": None,
                    "round": "prelim",
                })
            continue

        if round_value == "final":
            if row.final_best is not None:
                payload.append({
                    **base,
                    "score": row.final_best,
                    "rank": row.rank if is_rankable_score(row.final_best) else None,
                    "date": None,
                    "round": "final",
                })
            continue

        if round_value in {"semi", "one"}:
            continue

        if row.prelim_best is not None:
            payload.append({
                **base,
                "score": row.prelim_best,
                "rank": row.rank if is_rankable_score(row.prelim_best) else None,
                "date": None,
                "round": "prelim",
            })
        if row.final_best is not None:
            payload.append({
                **base,
                "score": row.final_best,
                "rank": row.rank if is_rankable_score(row.final_best) else None,
                "date": None,
                "round": "final",
            })

    return _normalize_results_payload(payload)


def list_jump_results(event_name: str, round_value: str | None, college: str | None, name: str | None, student_id: str | None, gender: str | None, db: Session):
    query = select(ResultJump).where(ResultJump.event == event_name)
    if college:
        query = query.where(ResultJump.college == college)
    if name:
        query = query.where(ResultJump.name == name)
    if student_id:
        query = query.where(ResultJump.student_id == student_id)
    if gender:
        athlete_query = select(Athlete.student_id).where(
            Athlete.event == event_name,
            Athlete.gender.in_(gender_filter_values(gender)),
        )
        student_ids = [row[0] for row in db.execute(athlete_query).all()]
        if not student_ids:
            return []
        query = query.where(ResultJump.student_id.in_(student_ids))
    if round_value:
        query = query.where(ResultJump.round == round_value)
    rows = db.scalars(query).all()

    athlete_map = {}
    if rows:
        athlete_query = select(Athlete).where(
            Athlete.student_id.in_([row.student_id for row in rows]),
            Athlete.event == event_name,
        )
        athlete_map = {
            (row.student_id, row.event): row.record_broken
            for row in db.scalars(athlete_query).all()
        }

    payload = [
        {
            "id": row.id,
            "event": row.event,
            "name": row.name,
            "student_id": row.student_id,
            "college": row.college,
            "score": row.score,
            "rank": row.rank if is_rankable_score(row.score) else None,
            "date": row.date,
            "round": row.round,
            "record_broken": athlete_map.get((row.student_id, row.event), False),
        }
        for row in rows
    ]
    return _normalize_results_payload(payload)


def list_track_results(
    event: Optional[str],
    gender: Optional[str],
    college: Optional[str],
    name: Optional[str],
    student_id: Optional[str],
    round_value: Optional[str],
    db: Session,
):
    query = select(Result)
    if event:
        query = query.where(Result.event == event)
    if college:
        query = query.where(Result.college == college)
    if name:
        query = query.where(Result.name == name)
    if student_id:
        query = query.where(Result.student_id == student_id)
    if gender:
        athlete_query = select(Athlete.student_id)
        if event:
            athlete_query = athlete_query.where(Athlete.event == event)
        athlete_query = athlete_query.where(Athlete.gender.in_(gender_filter_values(gender)))
        student_ids = [row[0] for row in db.execute(athlete_query).all()]
        if not student_ids:
            return []
        query = query.where(Result.student_id.in_(student_ids))
    if round_value:
        query = query.where(Result.round == round_value)
    rows = db.scalars(query).all()

    athlete_map = {}
    if rows:
        athlete_query = select(Athlete).where(
            Athlete.student_id.in_([row.student_id for row in rows])
        )
        if event:
            athlete_query = athlete_query.where(Athlete.event == event)
        athlete_map = {
            (row.student_id, row.event): row.record_broken
            for row in db.scalars(athlete_query).all()
        }

    payload = [
        {
            "id": row.id,
            "event": row.event,
            "name": row.name,
            "student_id": row.student_id,
            "college": row.college,
            "score": row.score,
            "rank": row.rank if is_rankable_score(row.score) else None,
            "date": row.date,
            "round": row.round,
            "record_broken": athlete_map.get((row.student_id, row.event), False),
        }
        for row in rows
    ]
    return _normalize_results_payload(payload)


def is_distance_event(event_name: str) -> bool:
    return score_higher_better(event_name)


def sort_scores(scores: List[float], event_name: str):



    reverse = is_distance_event(event_name)



    return sorted(scores, reverse=reverse)











def rank_athletes_by_score(rows: List[Dict], event_name: str):



    reverse = is_distance_event(event_name)



    rows_sorted = sorted(
        [row for row in rows if is_rankable_score(row.get("score"))],
        key=lambda r: r["score"],
        reverse=reverse,
    )



    previous_score = None
    previous_rank = 0

    for idx, row in enumerate(rows_sorted, start=1):
        score = row["score"]
        if previous_score is not None and math.isclose(score, previous_score, rel_tol=0.0, abs_tol=1e-9):
            row["rank"] = previous_rank
        else:
            row["rank"] = idx
            previous_rank = idx
            previous_score = score



    return rows_sorted


def is_rankable_score(score) -> bool:
    try:
        numeric = float(score)
    except (TypeError, ValueError):
        return False
    return math.isfinite(numeric) and numeric > 0


def _normalize_results_payload(items: List[Dict]) -> List[Dict]:
    valid_items = []
    for item in items:
        score = item.get("score")
        if not is_rankable_score(score):
            continue
        normalized = dict(item)
        normalized["score"] = float(score)
        valid_items.append(normalized)

    rank_maps: Dict[tuple, Dict] = {}
    for item in valid_items:
        key = (item.get("event"), item.get("round"))
        rank_maps.setdefault(key, [])
        rank_maps[key].append(
            {
                "id": item.get("id"),
                "score": item["score"],
            }
        )

    resolved_ranks: Dict[tuple, Dict] = {}
    for key, rows in rank_maps.items():
        ranked_rows = rank_athletes_by_score(rows, key[0])
        resolved_ranks[key] = {row["id"]: row["rank"] for row in ranked_rows}

    for item in valid_items:
        key = (item.get("event"), item.get("round"))
        item["rank"] = resolved_ranks.get(key, {}).get(item.get("id"))

    round_order = {"prelim": 0, "semi": 1, "final": 2, "one": 3}
    valid_items.sort(
        key=lambda item: (
            item.get("event") or "",
            round_order.get(item.get("round"), 9),
            item.get("rank") if item.get("rank") is not None else 9999,
            item.get("student_id") or "",
        )
    )
    return valid_items











RECORD_BROKEN_BONUS_POINTS = 9
RELAY_TEAM_POINTS = {1: 18, 2: 14, 3: 12, 4: 10, 5: 8, 6: 6, 7: 4, 8: 2}


def points_for_rank(rank: int) -> int:
    mapping = {1: 9, 2: 7, 3: 6, 4: 5, 5: 4, 6: 3, 7: 2, 8: 1}
    return mapping.get(rank, 0)


def points_for_relay_team_rank(rank: int) -> int:
    return RELAY_TEAM_POINTS.get(rank, 0)


def points_for_record_broken(record_broken: bool) -> int:
    return RECORD_BROKEN_BONUS_POINTS if record_broken else 0


def _points_for_export_result(row: Dict) -> int:
    try:
        rank_value = int(row.get("rank"))
    except (TypeError, ValueError):
        return 0
    if rank_value <= 0:
        return 0
    event_name = str(row.get("event") or "")
    base_points = (
        points_for_relay_team_rank(rank_value)
        if is_relay_event(event_name)
        else points_for_rank(rank_value)
    )
    return base_points + points_for_record_broken(bool(row.get("record_broken")))


def _format_export_time_score(score: Optional[float]) -> str:
    if not is_rankable_score(score):
        return ""
    numeric = float(score)
    hours = int(numeric // 3600)
    minutes = int((numeric - hours * 3600) // 60)
    seconds = round(numeric - hours * 3600 - minutes * 60, 2)
    second_text = f"{seconds:.2f}".rstrip("0").rstrip(".")
    if (hours > 0 or minutes > 0) and seconds < 10 and not second_text.startswith("0"):
        second_text = f"0{second_text}"
    if hours > 0:
        return f"{hours}:{minutes:02d}:{second_text}"
    if minutes > 0:
        return f"{minutes}:{second_text}"
    return second_text


def _format_export_distance_score(score: Optional[float]) -> str:
    if not is_rankable_score(score):
        return ""
    return f"{float(score):.3f}".rstrip("0").rstrip(".")


def _format_export_score(event_name: Optional[str], score: Optional[float]) -> str:
    normalized_event = (event_name or "").strip()
    if is_distance_event(normalized_event) or is_high_jump_event(normalized_event):
        return _format_export_distance_score(score)
    return _format_export_time_score(score)


def _format_export_rank_header(rank: Optional[int]) -> str:
    try:
        numeric_rank = int(rank)
    except (TypeError, ValueError):
        return ""
    if numeric_rank <= 0:
        return ""
    return f"第{_to_chinese_group_number(numeric_rank)}名"


def _build_ranked_results_sheet_stream(
    event: Optional[str],
    rows: List[Dict],
    round_value: Optional[str] = None,
    gender: Optional[str] = None,
) -> io.BytesIO:
    picked = list(_pick_export_rows_by_student(rows).values())
    ranked_rows = []
    for row in picked:
        score = row.get("score")
        if not is_rankable_score(score):
            continue
        ranked_rows.append(
            {
                "name": row.get("name", ""),
                "college": row.get("college", ""),
                "score": float(score),
                "display_score": _format_export_score(event, score),
                "points": _points_for_export_result(row),
            }
        )

    ranked_rows = rank_athletes_by_score(ranked_rows, event or "")
    ranked_rows = ranked_rows[:8]

    wb = Workbook()
    ws = wb.active
    ws.title = "results_sheet"

    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    title_align = Alignment(horizontal="center", vertical="center", wrap_text=False, shrink_to_fit=True)

    ws.merge_cells("A1:I1")
    title_parts = []
    gender_text = _gender_title_label(gender)
    if gender_text:
        title_parts.append(gender_text)
    event_text = (event or "成绩").strip() or "成绩"
    round_text = _round_display_text(round_value)
    title_parts.append(f"{event_text}{round_text}成绩单" if round_text else f"{event_text}成绩单")
    ws["A1"] = "    ".join(title_parts)
    ws["A1"].alignment = title_align
    ws.row_dimensions[1].height = 28

    headers = ["名次"] + [_format_export_rank_header(item.get("rank")) for item in ranked_rows]
    headers.extend([""] * (9 - len(headers)))
    labels = ["姓名", "单位", "成绩", "备注"]
    values_by_row = [
        [item.get("name", "") for item in ranked_rows],
        [item.get("college", "") for item in ranked_rows],
        [item.get("display_score", "") for item in ranked_rows],
        [item.get("points", 0) for item in ranked_rows],
    ]

    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=2, column=col_idx, value=header)
        cell.border = border
        cell.alignment = align

    for row_offset, label in enumerate(labels, start=3):
        row_index = row_offset
        ws.row_dimensions[row_index].height = 24
        label_cell = ws.cell(row=row_index, column=1, value=label)
        label_cell.border = border
        label_cell.alignment = align
        row_values = values_by_row[row_offset - 3]
        for col_offset in range(8):
            column_index = 2 + col_offset
            value = row_values[col_offset] if col_offset < len(row_values) else ""
            cell = ws.cell(row=row_index, column=column_index, value=value)
            cell.border = border
            cell.alignment = align

    ws.merge_cells("A7:I7")
    ws["A7"] = "总裁判：             裁判长：                裁判员：               记录员："
    ws["A7"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=False)
    ws.row_dimensions[7].height = 24

    ws.column_dimensions["A"].width = 10
    for col_idx in range(2, 10):
        ws.column_dimensions[get_column_letter(col_idx)].width = 14

    _apply_uniform_template_font(ws, 1, 7, 1, 9)
    ws.print_area = "A1:I7"

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output











def distribute_by_college(candidates: List[Athlete]) -> List[Athlete]:



    buckets: Dict[str, List[Athlete]] = {}



    for item in candidates:



        buckets.setdefault(item.college, []).append(item)



    for key in buckets:



        random.shuffle(buckets[key])



    ordered = []



    while any(buckets.values()):



        for college in list(buckets.keys()):



            if buckets[college]:



                ordered.append(buckets[college].pop())



    return ordered











def get_timestamp():



    global LAST_IMPORT_TS



    if LAST_IMPORT_TS:



        return LAST_IMPORT_TS



    return datetime.utcnow().strftime("%Y%m%d%H%M%S")



def build_results_template_filename(event: Optional[str], timestamp: str, gender: Optional[str] = None) -> str:
    label = (event or "results").strip() or "results"
    gender_label = _gender_filename_label(gender)
    if gender_label:
        label = f"{label}_{gender_label}"
    label = label.replace("*", "×")
    for ch in '\\/:?"<>|':
        label = label.replace(ch, "_")
    label = label.replace("\r", " ").replace("\n", " ").replace("\t", " ")
    return f"{label}_template_{timestamp}.xlsx"


def _format_export_date(value) -> str:
    if not value:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    return str(value)


def _round_display_text(round_value: Optional[str]) -> str:
    mapping = {
        "prelim": "\u9884\u8d5b",
        "semi": "\u534a\u51b3\u8d5b",
        "final": "\u51b3\u8d5b",
        "one": "\u4e00\u8f6e\u6b21",
    }
    return mapping.get((round_value or "").lower(), round_value or "")


def _gender_filename_label(gender: Optional[str]) -> str:
    raw = (gender or "").strip()
    text = raw.lower()
    if not text:
        return ""
    if "\u7537" in raw or text in {"male", "m", "man", "men"}:
        return "\u7537"
    if "\u5973" in raw or text in {"female", "f", "woman", "women"}:
        return "\u5973"
    if "\u6df7" in raw or text in {"mixed", "mix"}:
        return "\u6df7\u5408"
    return raw


def _gender_title_label(gender: Optional[str]) -> str:
    raw = (gender or "").strip()
    text = raw.lower()
    if not text:
        return ""
    if "\u7537" in raw or text in {"male", "m", "man", "men"}:
        return "\u7537\u5b50\u7ec4"
    if "\u5973" in raw or text in {"female", "f", "woman", "women"}:
        return "\u5973\u5b50\u7ec4"
    if "\u6df7" in raw or text in {"mixed", "mix"}:
        return "\u6df7\u5408\u7ec4"
    return raw


def _row_value(item, key: str, default=""):
    if isinstance(item, dict):
        return item.get(key, default)
    return getattr(item, key, default)


def _should_replace_result_payload(prev_row: Dict, new_row: Dict) -> bool:
    prev_priority = _round_priority_for_ranking(prev_row.get("round"))
    new_priority = _round_priority_for_ranking(new_row.get("round"))
    if new_priority != prev_priority:
        return new_priority > prev_priority
    prev_date = prev_row.get("date") or datetime.min
    new_date = new_row.get("date") or datetime.min
    return new_date >= prev_date


def _pick_export_rows_by_student(rows: List[Dict]) -> Dict[str, Dict]:
    picked: Dict[str, Dict] = {}
    for row in rows:
        student_key = str(row.get("student_id") or "").strip()
        if not student_key:
            continue
        prev = picked.get(student_key)
        if prev is None or _should_replace_result_payload(prev, row):
            picked[student_key] = row
    return picked


def _merge_template_athletes(athletes: List[Athlete], rows: List[Dict]) -> List:
    merged = list(athletes)
    existing = {str(_row_value(item, "student_id", "")).strip() for item in merged}
    for row in rows:
        student_key = str(row.get("student_id") or "").strip()
        if not student_key or student_key in existing:
            continue
        merged.append(
            {
                "student_id": student_key,
                "name": row.get("name", ""),
                "college": row.get("college", ""),
            }
        )
        existing.add(student_key)
    return merged


def _build_excel_stream_response(output: io.BytesIO, filename: str) -> StreamingResponse:
    safe_name = filename.encode("ascii", "ignore").decode("ascii") or "export.xlsx"
    encoded_name = urllib.parse.quote(filename)
    headers = {
        "Content-Disposition": f"attachment; filename={safe_name}; filename*=UTF-8''{encoded_name}"
    }
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers,
    )


def to_dict(obj):



    return {column.name: getattr(obj, column.name) for column in obj.__table__.columns}











@app.post("/auth/register")



def register(req: RegisterRequest, db: Session = Depends(get_db)):



    exists = db.scalar(select(User).where(User.email == req.email))



    if exists:



        raise HTTPException(status_code=400, detail="Email already registered")



    user = User(



        school=req.school,



        email=req.email,



        password_hash=hash_password(req.password),



        must_reset=False,



    )



    db.add(user)



    db.commit()



    return {"message": "registered"}











@app.post("/auth/login")



def login(req: LoginRequest, db: Session = Depends(get_db)):



    user = db.scalar(select(User).where(User.school == req.school))



    if not user or not verify_password(req.password, user.password_hash):



        raise HTTPException(status_code=401, detail="Invalid credentials")



    return {



        "message": "ok",



        "school": user.school,



        "email": user.email,



        "must_reset": user.must_reset,



    }











@app.post("/auth/reset-password")



def reset_password(req: ResetPasswordRequest, db: Session = Depends(get_db)):



    user = db.scalar(select(User).where(User.email == req.email))



    if not user:



        raise HTTPException(status_code=404, detail="Email not found")



    user.password_hash = hash_password(req.new_password)



    user.must_reset = False



    db.commit()



    return {"message": "password updated"}











@app.get("/athletes")



def list_athletes(
    college: Optional[str] = None,
    name: Optional[str] = None,
    student_id: Optional[str] = None,
    gender: Optional[str] = None,
    event: Optional[str] = None,
    group_name: Optional[str] = None,
    db: Session = Depends(get_db),
):
    query = select(Athlete)
    if college:
        query = query.where(Athlete.college == college)
    if name:
        query = query.where(Athlete.name == name)
    if student_id:
        query = query.where(Athlete.student_id == student_id)
    if gender:
        query = query.where(Athlete.gender.in_(gender_filter_values(gender)))
    if event:
        query = query.where(Athlete.event == event)
    if group_name:
        query = query.where(Athlete.group_name == group_name)

    rows = db.scalars(query).all()
    data = []
    for row in rows:
        item = to_dict(row)
        current_group = item.get("group_name") or ""
        latest_grouping = db.scalar(
            select(Grouping)
            .where(
                Grouping.student_id == row.student_id,
                Grouping.event == row.event,
            )
            .order_by(Grouping.id.desc())
        )
        # Prefer latest grouping label for display to avoid persisting historical mojibake text.
        if latest_grouping and latest_grouping.group_label:
            item["group_name"] = format_group_name(row.gender, latest_grouping.group_label)
        elif _looks_like_mojibake(current_group):
            item["group_name"] = ""
        data.append(item)
    return data


def _looks_like_mojibake(value: str) -> bool:
    if not value:
        return False
    if "??" in value:
        return True
    mojibake_marks = ("锟", "鈥", "闆", "鍒", "缁", "璇", "閫", "鎬", "鍙", "å", "æ", "ç", "é", "è", "ï", "�")
    hit = sum(value.count(mark) for mark in mojibake_marks)
    return hit >= 2


def query_groupings(
    db: Session,
    event: Optional[str] = None,
    gender: Optional[str] = None,
    round: Optional[str] = None,
):
    query = select(Grouping)
    if event:
        query = query.where(Grouping.event == event)
    if round:
        query = query.where(Grouping.round == round)
    if gender:
        athlete_query = select(Athlete.student_id)
        if event:
            athlete_query = athlete_query.where(Athlete.event == event)
        athlete_query = athlete_query.where(Athlete.gender.in_(gender_filter_values(gender)))
        student_ids = [row[0] for row in db.execute(athlete_query).all()]
        if not student_ids:
            return []
        query = query.where(Grouping.student_id.in_(student_ids))
    rows = db.scalars(query).all()
    return sorted(
        rows,
        key=lambda r: (
            r.event or "",
            r.round or "",
            r.group_label or "",
            r.lane if r.lane is not None else 999,
            r.student_id or "",
        ),
    )











@app.post("/athletes")



def create_athlete(req: AthleteCreate, db: Session = Depends(get_db)):



    athlete = Athlete(**req.model_dump())



    db.add(athlete)
    db.flush()
    sync_event_participant_counts(db)



    db.commit()



    db.refresh(athlete)



    return to_dict(athlete)











@app.put("/athletes/{athlete_id}")



def update_athlete(athlete_id: int, req: AthleteUpdate, db: Session = Depends(get_db)):



    athlete = db.get(Athlete, athlete_id)



    if not athlete:



        raise HTTPException(status_code=404, detail="Not found")



    for key, value in req.model_dump(exclude_unset=True).items():



        setattr(athlete, key, value)



    db.flush()
    sync_event_participant_counts(db)
    db.commit()



    return to_dict(athlete)











@app.delete("/athletes/{athlete_id}")



def delete_athlete(athlete_id: int, db: Session = Depends(get_db)):



    athlete = db.get(Athlete, athlete_id)



    if not athlete:



        raise HTTPException(status_code=404, detail="Not found")



    db.delete(athlete)
    db.flush()
    sync_event_participant_counts(db)



    db.commit()



    return {"message": "deleted"}


def sync_event_participant_counts(db: Session):
    counts: Dict[str, int] = {}
    for event_name in db.scalars(select(Athlete.event)).all():
        if not event_name:
            continue
        counts[event_name] = counts.get(event_name, 0) + 1

    existing_events = {row.name: row for row in db.scalars(select(Event)).all()}
    for event_name, event_row in existing_events.items():
        event_row.participant_count = counts.get(event_name, 0)

    for event_name, count in counts.items():
        if event_name not in existing_events:
            db.add(
                Event(
                    name=event_name,
                    participant_count=count,
                    best_record=None,
                    record_holder="",
                )
            )










ATHLETE_TEMPLATE_HEADERS = {
    "college": "学院",
    "name": "姓名",
    "student_id": "学号",
    "gender": "性别",
    "event": "项目",
    "group_name": "组别",
    "score": "成绩",
    "rank": "排名",
    "points": "积分",
    "phone": "电话",
}

ATHLETE_TEMPLATE_ORDER = [
    "college",
    "name",
    "student_id",
    "gender",
    "event",
    "group_name",
    "score",
    "rank",
    "points",
    "phone",
]


@app.post("/athletes/import")
def import_athletes(file: UploadFile = File(...), db: Session = Depends(get_db)):
    import re

    try:
        df = pd.read_excel(file.file)
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Unable to read Excel file: {exc}")

    if df.empty:
        raise HTTPException(status_code=400, detail="Import file is empty")

    def normalize(name: str) -> str:
        value = str(name or "").strip().lower()
        value = re.sub(r"[\s\-_()（）【】\[\]:：]+", "", value)
        return value

    aliases = {
        "college": ["学院", "院系", "college", "collegename"],
        "name": ["姓名", "运动员姓名", "name", "studentname"],
        "student_id": ["学号", "学号id", "studentid", "student_no", "id"],
        "gender": ["性别", "gender", "sex"],
        "event": ["项目", "比赛项目", "event", "eventname"],
        "group_name": ["组别", "分组", "group", "groupname"],
        "score": ["成绩", "score", "result"],
        "rank": ["排名", "名次", "rank"],
        "points": ["积分", "得分", "points", "point"],
        "phone": ["电话", "手机号", "联系方式", "phone", "mobile"],
    }

    normalized_columns = {normalize(col): col for col in list(df.columns)}
    column_map = {}
    for field, names in aliases.items():
        for alias in names:
            key = normalize(alias)
            if key in normalized_columns:
                column_map[field] = normalized_columns[key]
                break

    required_fields = ["college", "name", "student_id", "gender", "event"]
    if any(field not in column_map for field in required_fields):
        cols = list(df.columns)
        for idx, field in enumerate(ATHLETE_TEMPLATE_ORDER):
            if idx < len(cols) and field not in column_map:
                column_map[field] = cols[idx]

    missing = [field for field in required_fields if field not in column_map]
    if missing:
        uploaded = ", ".join(str(c) for c in list(df.columns))
        raise HTTPException(
            status_code=400,
            detail=f"Import failed: required columns missing (college/name/student_id/gender/event). Current columns: {uploaded}",
        )

    def cell(row, field):
        col = column_map.get(field)
        if not col:
            return None
        return row[col]

    def as_text(value):
        if value is None or (isinstance(value, float) and pd.isna(value)):
            return ""
        return str(value).strip()

    def as_int(value):
        if value is None or pd.isna(value):
            return None
        try:
            return int(float(value))
        except Exception:
            return None

    def as_float(value):
        if value is None or pd.isna(value):
            return None
        try:
            return float(value)
        except Exception:
            return None

    prepared: Dict[tuple, Dict] = {}
    invalid_rows: List[str] = []
    merged_duplicates = 0

    for excel_rownum, row in enumerate(df.iterrows(), start=2):
        _, row_data = row
        college = as_text(cell(row_data, "college"))
        name = as_text(cell(row_data, "name"))
        student_id = as_text(cell(row_data, "student_id"))
        gender = normalize_gender_value(as_text(cell(row_data, "gender")))
        event_name = as_text(cell(row_data, "event"))
        group_name = as_text(cell(row_data, "group_name")) or None
        score = as_float(cell(row_data, "score"))
        rank = as_int(cell(row_data, "rank"))
        points = as_int(cell(row_data, "points"))
        phone = as_text(cell(row_data, "phone")) or None

        if not any([college, name, student_id, gender, event_name, group_name, phone, score, rank, points]):
            continue

        if not college or not name or not gender or not event_name:
            invalid_rows.append(f"Row {excel_rownum}: missing required fields among college/name/gender/event")
            continue

        relay_event = is_relay_event(event_name)
        if not relay_event and not student_id:
            invalid_rows.append(f"Row {excel_rownum}: individual event missing student_id")
            continue

        logical_identity = f"{name}|{gender}" if relay_event else student_id
        if not logical_identity:
            invalid_rows.append(f"Row {excel_rownum}: missing unique identity value")
            continue

        storage_student_id = student_id or build_relay_storage_student_id(event_name, name, gender)
        logical_key = (normalize_event_name(event_name), logical_identity.strip())
        candidate = {
            "college": college,
            "name": name,
            "student_id": storage_student_id,
            "gender": gender,
            "event": event_name,
            "group_name": group_name,
            "score": score,
            "rank": rank,
            "points": points,
            "phone": phone,
            "_rownum": excel_rownum,
            "_relay": relay_event,
            "_identity": logical_identity.strip(),
        }

        existing = prepared.get(logical_key)
        if existing is None:
            prepared[logical_key] = candidate
            continue

        conflict_fields = ["college", "name", "gender", "event"]
        if not relay_event:
            conflict_fields.append("student_id")
        conflicts = [
            field
            for field in conflict_fields
            if existing.get(field) not in {None, ""}
            and candidate.get(field) not in {None, ""}
            and existing.get(field) != candidate.get(field)
        ]
        if conflicts:
            key_label = "name+gender+event" if relay_event else "student_id+event"
            raise HTTPException(
                status_code=400,
                detail=(
                    f"Import failed: duplicate {key_label} conflict at rows "
                    f"{existing['_rownum']} and {excel_rownum}: {logical_identity} / {event_name}"
                ),
            )

        for optional_field in ["group_name", "score", "rank", "points", "phone"]:
            if existing.get(optional_field) in {None, ""} and candidate.get(optional_field) not in {None, ""}:
                existing[optional_field] = candidate[optional_field]
        merged_duplicates += 1

    if invalid_rows:
        preview = " | ".join(invalid_rows[:10])
        if len(invalid_rows) > 10:
            preview += f" | and {len(invalid_rows) - 10} more"
        raise HTTPException(status_code=400, detail=f"Import failed: {preview}")

    db.execute(delete(Athlete))
    db.execute(delete(Grouping))
    db.execute(delete(Result))
    db.execute(delete(ResultDistance))
    db.execute(delete(ResultJump))
    db.execute(delete(PersonalRanking))
    db.execute(delete(CollegeRanking))
    db.commit()

    global LAST_IMPORT_TS
    LAST_IMPORT_TS = datetime.utcnow().strftime("%Y%m%d%H%M%S")

    inserted = 0
    for item in prepared.values():
        athlete = Athlete(
            college=item["college"],
            name=item["name"],
            student_id=item["student_id"],
            gender=item["gender"],
            event=item["event"],
            group_name=item["group_name"],
            score=item["score"],
            rank=item["rank"],
            points=item["points"],
            phone=item["phone"],
        )
        db.add(athlete)
        inserted += 1

    db.commit()

    sync_event_participant_counts(db)
    db.commit()

    suffix = f", merged {merged_duplicates} duplicate rows" if merged_duplicates else ""
    return {"message": f"imported {inserted} rows{suffix}"}


@app.get("/athletes/export")
def export_athletes(db: Session = Depends(get_db)):
    rows = db.scalars(select(Athlete)).all()
    data = [
        {
            ATHLETE_TEMPLATE_HEADERS["college"]: r.college,
            ATHLETE_TEMPLATE_HEADERS["name"]: r.name,
            ATHLETE_TEMPLATE_HEADERS["student_id"]: r.student_id,
            ATHLETE_TEMPLATE_HEADERS["gender"]: r.gender,
            ATHLETE_TEMPLATE_HEADERS["event"]: r.event,
            ATHLETE_TEMPLATE_HEADERS["group_name"]: r.group_name or "",
            ATHLETE_TEMPLATE_HEADERS["score"]: r.score if r.score is not None else "",
            ATHLETE_TEMPLATE_HEADERS["rank"]: r.rank if r.rank is not None else "",
            ATHLETE_TEMPLATE_HEADERS["points"]: r.points if r.points is not None else "",
            ATHLETE_TEMPLATE_HEADERS["phone"]: r.phone or "",
        }
        for r in rows
    ]
    timestamp = get_timestamp()
    return excel_response(data, f"athletes_{timestamp}.xlsx")


@app.get("/athletes/template")
def export_athletes_template():
    data = [{ATHLETE_TEMPLATE_HEADERS[field]: "" for field in ATHLETE_TEMPLATE_ORDER}]
    timestamp = get_timestamp()
    return excel_response(data, f"athletes_template_{timestamp}.xlsx")


@app.get("/groups")



def list_groups(



    event: Optional[str] = None,



    gender: Optional[str] = None,



    round: Optional[str] = None,



    db: Session = Depends(get_db),



):



    query = select(Grouping)



    if event:



        query = query.where(Grouping.event == event)



    if round:



        query = query.where(Grouping.round == round)



    if gender:



        athlete_query = select(Athlete.student_id)



        if event:



            athlete_query = athlete_query.where(Athlete.event == event)



        athlete_query = athlete_query.where(Athlete.gender == gender)



        student_ids = [row[0] for row in db.execute(athlete_query).all()]



        if not student_ids:



            return []



        query = query.where(Grouping.student_id.in_(student_ids))



    return [to_dict(row) for row in db.scalars(query).all()]













@app.post("/groups/confirm")
def confirm_groups(req: GroupConfirmRequest, db: Session = Depends(get_db)):
    delete_query = delete(Grouping).where(
        Grouping.event == req.event,
        Grouping.round == req.round,
    )
    target_gender = (req.gender or "").strip()
    if target_gender:
        student_ids = [
            row[0]
            for row in db.execute(
                select(Athlete.student_id).where(
                    Athlete.event == req.event,
                    Athlete.gender.in_(gender_filter_values(target_gender)),
                )
            ).all()
        ]
        if student_ids:
            delete_query = delete_query.where(Grouping.student_id.in_(student_ids))
        else:
            delete_query = None
    if delete_query is not None:
        db.execute(delete_query)
    for item in req.groups:
        grouping = Grouping(**item.model_dump())
        db.add(grouping)
        athlete = db.scalar(
            select(Athlete).where(
                Athlete.student_id == item.student_id,
                Athlete.event == req.event,
            )
        )
        if athlete:
            athlete.group_name = format_group_name(athlete.gender, item.group_label)
    db.commit()
    return {"message": "saved"}

@app.get("/events/list")



def list_events(db: Session = Depends(get_db)):



    return [to_dict(row) for row in db.scalars(select(Event)).all()]











@app.get("/events/options")



def list_event_options(db: Session = Depends(get_db)):



    rows = db.execute(select(distinct(Athlete.event))).all()



    return [row[0] for row in rows if row[0]]











@app.get("/athletes/by-event")



def list_athletes_by_event(event: str, db: Session = Depends(get_db)):



    query = select(Athlete).where(Athlete.event == event)



    return [to_dict(row) for row in db.scalars(query).all()]











@app.get("/groups/candidates")
def list_group_candidates(
    event: str,
    gender: str,
    round: str,
    db: Session = Depends(get_db),
):
    eligible_athletes = db.scalars(
        select(Athlete).where(
            Athlete.event == event,
            Athlete.gender.in_(gender_filter_values(gender)),
        )
    ).all()
    if round != "final":
        return [to_dict(row) for row in eligible_athletes]

    eligible_ids = {row.student_id for row in eligible_athletes}
    if not eligible_ids:
        return []

    kind = classify_event_table(event)
    scored: List[Dict] = []

    if kind == "distance" or has_distance_records(event, db):
        distance_rows = db.scalars(
            select(ResultDistance).where(
                ResultDistance.event == event,
                ResultDistance.prelim_best.isnot(None),
            )
        ).all()
        scored = [
            {"student_id": r.student_id, "score": r.prelim_best}
            for r in distance_rows
            if r.prelim_best is not None and r.student_id in eligible_ids
        ]
    elif kind == "jump" or has_jump_records(event, db):
        jump_rows = db.scalars(
            select(ResultJump).where(
                ResultJump.event == event,
                ResultJump.round.in_(["prelim", "one"]),
            )
        ).all()
        scored = [
            {"student_id": r.student_id, "score": r.score}
            for r in jump_rows
            if r.score is not None and r.student_id in eligible_ids
        ]
    else:
        result_rows = db.scalars(
            select(Result).where(
                Result.event == event,
                Result.round.in_(["prelim", "one"]),
            )
        ).all()
        scored = [
            {"student_id": r.student_id, "score": r.score}
            for r in result_rows
            if r.score is not None and r.student_id in eligible_ids
        ]

    if not scored:
        prelim_groups = db.scalars(
            select(Grouping).where(
                Grouping.event == event,
                Grouping.round.in_(["prelim", "one"]),
            )
        ).all()
        for row in prelim_groups:
            if row.student_id not in eligible_ids:
                continue
            score = row.prelim_score if row.round in {"prelim", "one"} else None
            if score is None:
                continue
            scored.append({"student_id": row.student_id, "score": score})

    ranked_ids = [r["student_id"] for r in rank_athletes_by_score(scored, event)[:8]]
    if not ranked_ids:
        raise HTTPException(
            status_code=404,
            detail="未找到可用于决赛分组的预赛成绩，请先录入预赛成绩",
        )

    athlete_map = {row.student_id: row for row in eligible_athletes}
    athletes = [athlete_map[sid] for sid in ranked_ids if sid in athlete_map]
    return [to_dict(row) for row in athletes]


@app.get("/results/list")



def list_results(
    event: Optional[str] = None,
    gender: Optional[str] = None,
    college: Optional[str] = None,
    name: Optional[str] = None,
    student_id: Optional[str] = None,
    round: Optional[str] = None,
    db: Session = Depends(get_db),
):
    if event:
        kind = classify_event_table(event)
        if kind == "distance" or has_distance_records(event, db):
            return list_distance_results(event, round, college, name, student_id, gender, db)
        if kind == "jump" or has_jump_records(event, db):
            return list_jump_results(event, round, college, name, student_id, gender, db)
        return list_track_results(event, gender, college, name, student_id, round, db)

    merged = []
    merged.extend(list_track_results(None, gender, college, name, student_id, round, db))

    distance_events = [row[0] for row in db.execute(select(distinct(ResultDistance.event))).all() if row[0]]
    for event_name in distance_events:
        merged.extend(list_distance_results(event_name, round, college, name, student_id, gender, db))

    jump_events = [row[0] for row in db.execute(select(distinct(ResultJump.event))).all() if row[0]]
    for event_name in jump_events:
        merged.extend(list_jump_results(event_name, round, college, name, student_id, gender, db))

    round_order = {"prelim": 0, "semi": 1, "final": 2, "one": 3}
    merged.sort(
        key=lambda item: (
            item.get("event") or "",
            round_order.get(item.get("round"), 9),
            item.get("rank") if item.get("rank") is not None else 9999,
            item.get("student_id") or "",
        )
    )
    return merged


def _update_event_record(event_name: str, athlete: Athlete, score: float, higher_better: bool, db: Session):
    event_row = db.scalar(select(Event).where(Event.name == event_name))
    if not event_row:
        event_row = Event(name=event_name, participant_count=0, best_record=None, record_holder=None)
        db.add(event_row)

    if event_row.best_record is None:
        should_update = True
    else:
        should_update = score > event_row.best_record if higher_better else score < event_row.best_record

    if should_update:
        event_row.best_record = score
        event_row.record_holder = athlete.name
        athlete.record_broken = True


def save_track_result(req: ResultSubmitRequest, db: Session):
    athlete = db.scalar(
        select(Athlete).where(
            Athlete.student_id == req.student_id,
            Athlete.event == req.event,
        )
    )
    if not athlete:
        raise HTTPException(status_code=404, detail="Athlete not found")

    round_value = req.round or "final"
    row = db.scalar(
        select(Result).where(
            Result.student_id == req.student_id,
            Result.event == req.event,
            Result.round == round_value,
        )
    )
    if not row:
        row = Result(
            event=req.event,
            name=athlete.name,
            student_id=athlete.student_id,
            college=athlete.college,
            score=req.score,
            rank=0,
            date=req.date or datetime.utcnow(),
            round=round_value,
        )
        db.add(row)
    else:
        row.name = athlete.name
        row.college = athlete.college
        row.score = req.score
        row.date = req.date or datetime.utcnow()

    db.flush()

    scoped_rows = db.scalars(
        select(Result).where(
            Result.event == req.event,
            Result.round == round_value,
        )
    ).all()
    scored = [
        {"student_id": item.student_id, "score": item.score}
        for item in scoped_rows
        if is_rankable_score(item.score)
    ]
    ranks = {r["student_id"]: r["rank"] for r in rank_athletes_by_score(scored, req.event)}
    for item in scoped_rows:
        item.rank = ranks.get(item.student_id, 0)

    if round_value in {"final", "one"}:
        athlete.score = req.score
        athlete.rank = ranks.get(req.student_id)
        if is_rankable_score(req.score):
            _update_event_record(req.event, athlete, req.score, higher_better=False, db=db)

    db.commit()
    rebuild_personal_rankings(db)
    rebuild_college_rankings(db)
    return {"message": "saved"}


def save_jump_result(req: ResultSubmitRequest, db: Session):
    athlete = db.scalar(
        select(Athlete).where(
            Athlete.student_id == req.student_id,
            Athlete.event == req.event,
        )
    )
    if not athlete:
        raise HTTPException(status_code=404, detail="Athlete not found")

    round_value = req.round or "final"
    row = db.scalar(
        select(ResultJump).where(
            ResultJump.student_id == req.student_id,
            ResultJump.event == req.event,
            ResultJump.round == round_value,
        )
    )
    if not row:
        row = ResultJump(
            event=req.event,
            name=athlete.name,
            student_id=athlete.student_id,
            college=athlete.college,
            score=req.score,
            rank=0,
            date=req.date or datetime.utcnow(),
            round=round_value,
        )
        db.add(row)
    else:
        row.name = athlete.name
        row.college = athlete.college
        row.score = req.score
        row.date = req.date or datetime.utcnow()

    db.flush()

    scoped_rows = db.scalars(
        select(ResultJump).where(
            ResultJump.event == req.event,
            ResultJump.round == round_value,
        )
    ).all()
    scored = [
        {"student_id": item.student_id, "score": item.score}
        for item in scoped_rows
        if is_rankable_score(item.score)
    ]
    ranks = {r["student_id"]: r["rank"] for r in rank_athletes_by_score(scored, req.event)}
    for item in scoped_rows:
        item.rank = ranks.get(item.student_id)

    if round_value in {"final", "one"}:
        athlete.score = req.score
        athlete.rank = ranks.get(req.student_id)
        if is_rankable_score(req.score):
            _update_event_record(req.event, athlete, req.score, higher_better=True, db=db)

    db.commit()
    rebuild_personal_rankings(db)
    rebuild_college_rankings(db)
    return {"message": "saved"}


@app.post("/results/submit")
def submit_results(req: ResultSubmitRequest, db: Session = Depends(get_db)):
    kind = classify_event_table(req.event)
    if kind == "distance" or has_distance_records(req.event, db):
        return save_distance_result(req, db)
    if kind == "jump" or has_jump_records(req.event, db):
        return save_jump_result(req, db)
    return save_track_result(req, db)


@app.get("/results-distance/list")



def list_results_distance(



    event: str,



    name: Optional[str] = None,



    student_id: Optional[str] = None,



    round: Optional[str] = None,



    db: Session = Depends(get_db),



):



    return list_distance_results(event, round, name, student_id, db)













def save_distance_result(req: ResultSubmitRequest, db: Session):
    if not getattr(req, "attempt", None):
        raise HTTPException(status_code=400, detail="attempt is required for field events")
    athlete = db.scalar(
        select(Athlete).where(
            Athlete.student_id == req.student_id,
            Athlete.event == req.event,
        )
    )
    if not athlete:
        raise HTTPException(status_code=404, detail="Athlete not found")
    row = db.scalar(
        select(ResultDistance).where(
            ResultDistance.student_id == req.student_id,
            ResultDistance.event == req.event,
        )
    )
    if not row:
        row = ResultDistance(
            event=req.event,
            name=athlete.name,
            student_id=athlete.student_id,
            college=athlete.college,
        )
        db.add(row)
    round_value = req.round
    if round_value == "one":
        round_value = "final"
    attempt = int(req.attempt)
    if attempt not in {1, 2, 3}:
        raise HTTPException(status_code=400, detail="attempt must be 1-3")
    if round_value == "prelim":
        setattr(row, f"prelim_attempt{attempt}", req.score)
        row.prelim_best = best_distance(row.prelim_attempt1, row.prelim_attempt2, row.prelim_attempt3)
    else:
        setattr(row, f"final_attempt{attempt}", req.score)
        row.final_best = best_distance(row.final_attempt1, row.final_attempt2, row.final_attempt3)
    rank_distance_event(req.event, round_value, db)
    if round_value == "final":
        update_final_results_distance(req.event, db, req.date)
        best = row.final_best
        if is_rankable_score(best):
            athlete.score = best
            athlete.rank = row.rank
            event_row = db.scalar(select(Event).where(Event.name == req.event))
            if not event_row:
                event_row = Event(name=req.event, participant_count=0, best_record=None, record_holder=None)
                db.add(event_row)
            should_update = event_row.best_record is None or best > event_row.best_record
            if should_update:
                event_row.best_record = best
                event_row.record_holder = athlete.name if athlete else ""
                athlete.record_broken = True
    db.commit()
    rebuild_personal_rankings(db)
    rebuild_college_rankings(db)
    return {"message": "saved"}

@app.post("/results-distance/submit")
def submit_results_distance(req: ResultSubmitRequest, db: Session = Depends(get_db)):
    return save_distance_result(req, db)


@app.get("/export/results-distance")



def export_results_distance(event: str, db: Session = Depends(get_db)):
    rows = db.scalars(select(ResultDistance).where(ResultDistance.event == event)).all()
    data = [
        {
            "\u9879\u76ee": r.event,
            "\u59d3\u540d": r.name,
            "\u5b66\u53f7": r.student_id,
            "\u5b66\u9662": r.college,
            "\u9884\u8d5b\u7b2c\u4e00\u6b21": r.prelim_attempt1 or "",
            "\u9884\u8d5b\u7b2c\u4e8c\u6b21": r.prelim_attempt2 or "",
            "\u9884\u8d5b\u7b2c\u4e09\u6b21": r.prelim_attempt3 or "",
            "\u9884\u8d5b\u6700\u4f73": r.prelim_best or "",
            "\u51b3\u8d5b\u7b2c\u4e00\u6b21": r.final_attempt1 or "",
            "\u51b3\u8d5b\u7b2c\u4e8c\u6b21": r.final_attempt2 or "",
            "\u51b3\u8d5b\u7b2c\u4e09\u6b21": r.final_attempt3 or "",
            "\u51b3\u8d5b\u6700\u4f73": r.final_best or "",
            "\u6392\u540d": r.rank or "",
        }
        for r in rows
    ]
    timestamp = get_timestamp()
    return excel_response(data, f"results_distance_{event}_{timestamp}.xlsx")



@app.get("/export/groups")
def export_groups(
    event: Optional[str] = None,
    gender: Optional[str] = None,
    round: Optional[str] = None,
    db: Session = Depends(get_db),
):
    rows = query_groupings(db=db, event=event, gender=gender, round=round)
    data = [
        {
            "event": r.event,
            "round": r.round,
            "group": r.group_label,
            "lane": r.lane if r.lane is not None else "",
            "college": r.college,
            "student_id": r.student_id,
            "name": r.name,
            "prelim_score": r.prelim_score if r.prelim_score is not None else "",
            "semi_score": r.semi_score if r.semi_score is not None else "",
            "final_score": r.final_score if r.final_score is not None else "",
        }
        for r in rows
    ]
    timestamp = get_timestamp()
    return excel_response(data, f"groups_{timestamp}.xlsx")


@app.get("/export/groups-template")
def export_groups_template(
    event: Optional[str] = None,
    gender: Optional[str] = None,
    round: Optional[str] = None,
    db: Session = Depends(get_db),
):
    rows = query_groupings(db=db, event=event, gender=gender, round=round)
    if rows:
        data = [
            {
                "event": r.event,
                "round": r.round,
                "group": r.group_label,
                "lane": r.lane if r.lane is not None else "",
                "college": r.college,
                "student_id": r.student_id,
                "name": r.name,
                "score": "",
                "rank": "",
                "notes": "",
            }
            for r in rows
        ]
    else:
        data = [
            {
                "event": event or "",
                "round": round or "prelim",
                "group": "A",
                "lane": "",
                "college": "",
                "student_id": "",
                "name": "",
                "score": "",
                "rank": "",
                "notes": "",
            }
        ]
    timestamp = get_timestamp()
    return excel_response(data, f"groups_template_{timestamp}.xlsx")


@app.get("/export/events")
def export_events(db: Session = Depends(get_db)):
    rows = db.scalars(select(Event)).all()
    data = [
        {
            "\u9879\u76ee": r.name,
            "\u53c2\u8d5b\u4eba\u6570": r.participant_count,
            "\u5386\u53f2\u6700\u4f73": r.best_record if r.best_record is not None else "",
            "\u7eaa\u5f55\u4fdd\u6301\u8005": r.record_holder or "",
        }
        for r in rows
    ]
    timestamp = get_timestamp()
    return excel_response(data, f"events_{timestamp}.xlsx")











@app.get("/export/results")
def export_results(
    event: Optional[str] = None,
    gender: Optional[str] = None,
    college: Optional[str] = None,
    name: Optional[str] = None,
    student_id: Optional[str] = None,
    round: Optional[str] = None,
    db: Session = Depends(get_db),
):
    if not event:
        raise HTTPException(status_code=400, detail="请先选择项目后再导出")

    rows = list_results(event=event, gender=gender, college=college, name=name, student_id=student_id, round=round, db=db)
    filename = build_results_template_filename(event, get_timestamp(), gender)
    output = _build_ranked_results_sheet_stream(
        event,
        rows,
        round_value=round,
        gender=gender,
    )
    return _build_excel_stream_response(output, filename)


def _list_template_athletes(
    event: Optional[str],
    db: Session,
    gender: Optional[str] = None,
    college: Optional[str] = None,
    name: Optional[str] = None,
    student_id: Optional[str] = None,
) -> List[Athlete]:
    query = select(Athlete)
    if event:
        query = query.where(Athlete.event == event)
    if gender:
        query = query.where(Athlete.gender.in_(gender_filter_values(gender)))
    if college:
        query = query.where(Athlete.college == college)
    if name:
        query = query.where(Athlete.name == name)
    if student_id:
        query = query.where(Athlete.student_id == student_id)
    query = query.order_by(Athlete.college.asc(), Athlete.name.asc(), Athlete.student_id.asc())
    return db.scalars(query).all()


def _list_template_groupings(
    event: Optional[str],
    round_value: Optional[str],
    db: Session,
    gender: Optional[str] = None,
) -> tuple[List[Grouping], Optional[str]]:
    query = select(Grouping)
    if event:
        query = query.where(Grouping.event == event)

    rows = db.scalars(query).all()
    if not rows:
        return [], None

    if gender:
        athlete_query = select(Athlete.student_id)
        if event:
            athlete_query = athlete_query.where(Athlete.event == event)
        athlete_query = athlete_query.where(Athlete.gender.in_(gender_filter_values(gender)))
        student_ids = {row[0] for row in db.execute(athlete_query).all()}
        if not student_ids:
            return [], None
        rows = [row for row in rows if row.student_id in student_ids]
        if not rows:
            return [], None

    selected_round = round_value
    if round_value:
        exact_rows = [row for row in rows if row.round == round_value]
        if exact_rows:
            rows = exact_rows
        else:
            selected_round = None

    if not selected_round:
        round_priority = {"prelim": 1, "semi": 2, "one": 3, "final": 4}
        best_round = max(rows, key=lambda row: round_priority.get((row.round or "").lower(), 0)).round
        rows = [row for row in rows if row.round == best_round]
        selected_round = best_round

    rows.sort(
        key=lambda row: (
            _group_label_sort_value(row.group_label or ""),
            row.lane if row.lane is not None else 999,
            str(row.name or ""),
            str(row.student_id or ""),
        )
    )
    return rows, selected_round


TEMPLATE_FONT_NAME = "SimSun"
TEMPLATE_FONT_SIZE = 12


def _apply_uniform_template_font(ws, min_row: int, max_row: int, min_col: int, max_col: int):
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            font = copy(cell.font)
            font.name = TEMPLATE_FONT_NAME
            font.size = TEMPLATE_FONT_SIZE
            cell.font = font


def _build_jump_results_template_stream(
    athletes: List[Athlete],
    result_by_student: Optional[Dict[str, Dict]] = None,
) -> io.BytesIO:
    result_by_student = result_by_student or {}
    template_path = os.path.join(os.path.dirname(__file__), "templates", "results_jump_style_template.xlsx")
    using_template = os.path.exists(template_path)
    if using_template:
        wb = load_workbook(template_path)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "jump_template"

    required_merges = ["A1:A2", "B1:B2", "C1:C2", "D1:AG1", "AH1:AH2", "AI1:AI2", "AJ1:AJ2"]
    existing_merges = {str(rng) for rng in ws.merged_cells.ranges}
    for merge_range in required_merges:
        if merge_range not in existing_merges:
            ws.merge_cells(merge_range)

    for col in range(4, 34, 3):
        start_col = get_column_letter(col)
        end_col = get_column_letter(col + 2)
        merge_range = f"{start_col}2:{end_col}2"
        if merge_range not in existing_merges:
            ws.merge_cells(merge_range)

    ws["A1"] = "\u5e8f\u53f7"
    ws["B1"] = "\u59d3\u540d"
    ws["C1"] = "\u5b66\u9662"
    ws["D1"] = "\u9ad8\u5ea6"
    ws["AH1"] = "\u6210\u7ee9"
    ws["AI1"] = "\u540d\u8bcd"
    ws["AJ1"] = "\u5907\u6ce8"

    for idx, athlete in enumerate(athletes, start=1):
        row_idx = idx + 2
        student_key = str(_row_value(athlete, "student_id", "")).strip()
        result_item = result_by_student.get(student_key, {})
        ws.cell(row=row_idx, column=1, value=idx)
        ws.cell(row=row_idx, column=2, value=_row_value(athlete, "name", ""))
        ws.cell(row=row_idx, column=3, value=_row_value(athlete, "college", ""))
        ws.cell(row=row_idx, column=34, value=result_item.get("score") if result_item.get("score") is not None else "")
        ws.cell(row=row_idx, column=35, value=result_item.get("rank") if result_item.get("rank") is not None else "")
        ws.cell(row=row_idx, column=36, value="")

    if not using_template:
        thin = Side(style="thin", color="000000")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        align = Alignment(horizontal="center", vertical="center", wrap_text=True)

        max_rows = max(77, len(athletes) + 2)
        for r in range(1, max_rows + 1):
            for c in range(1, 37):
                cell = ws.cell(row=r, column=c)
                cell.border = border
                cell.alignment = align

        ws.row_dimensions[1].height = 24
        ws.row_dimensions[2].height = 22
        ws.column_dimensions["A"].width = 8
        ws.column_dimensions["B"].width = 12
        ws.column_dimensions["C"].width = 16
        for c in range(4, 34):
            ws.column_dimensions[get_column_letter(c)].width = 4.6
        ws.column_dimensions["AH"].width = 10
        ws.column_dimensions["AI"].width = 10
        ws.column_dimensions["AJ"].width = 18
        ws.freeze_panes = "A3"
        ws.print_area = f"A1:AJ{max_rows}"

    max_rows = max(ws.max_row, len(athletes) + 2, 77 if not using_template else 0)
    _apply_uniform_template_font(ws, 1, max_rows, 1, 36)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def _build_field_results_template_stream(
    athletes: List[Athlete],
    result_by_student: Optional[Dict[str, Dict]] = None,
    round_value: Optional[str] = None,
) -> io.BytesIO:
    result_by_student = result_by_student or {}
    template_path = os.path.join(
        os.path.dirname(__file__),
        "templates",
        "results_standing_long_jump_style_template.xlsx",
    )
    using_template = os.path.exists(template_path)
    if using_template:
        wb = load_workbook(template_path)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "field_template"

    required_merges = [
        "A1:A2",
        "B1:B2",
        "C1:C2",
        "D1:F1",
        "G1:G2",
        "H1:J1",
        "K1:K2",
        "L1:L2",
        "M1:M2",
    ]
    existing_merges = {str(rng) for rng in ws.merged_cells.ranges}
    for merge_range in required_merges:
        if merge_range not in existing_merges:
            ws.merge_cells(merge_range)

    ws["A1"] = "\u53f7\u7801"
    ws["B1"] = "\u59d3\u540d"
    ws["C1"] = "\u5b66\u9662"
    ws["D1"] = "\u9884\u8d5b"
    ws["G1"] = "\u9884\u8d5b\u6700\u4f18"
    ws["H1"] = "\u51b3\u8d5b"
    ws["K1"] = "\u51b3\u8d5b\u6700\u4f18"
    ws["L1"] = "\u5907\u6ce8"
    ws["M1"] = "\u540d\u8bcd"
    ws["D2"] = "\u7b2c\u4e00\u6b21"
    ws["E2"] = "\u7b2c\u4e8c\u6b21"
    ws["F2"] = "\u7b2c\u4e09\u6b21"
    ws["H2"] = "\u7b2c\u4e00\u6b21"
    ws["I2"] = "\u7b2c\u4e8c\u6b21"
    ws["J2"] = "\u7b2c\u4e09\u6b21"

    show_prelim = round_value in {None, "prelim"}
    show_final = round_value in {None, "final"}

    for idx, athlete in enumerate(athletes, start=1):
        row_idx = idx + 2
        student_key = str(_row_value(athlete, "student_id", "")).strip()
        result_item = result_by_student.get(student_key, {})
        ws.cell(row=row_idx, column=1, value=idx)
        ws.cell(row=row_idx, column=2, value=_row_value(athlete, "name", ""))
        ws.cell(row=row_idx, column=3, value=_row_value(athlete, "college", ""))

        prelim_values = [
            result_item.get("prelim_attempt1"),
            result_item.get("prelim_attempt2"),
            result_item.get("prelim_attempt3"),
            result_item.get("prelim_best"),
        ]
        final_values = [
            result_item.get("final_attempt1"),
            result_item.get("final_attempt2"),
            result_item.get("final_attempt3"),
            result_item.get("final_best"),
        ]
        if show_prelim:
            for offset, value in enumerate(prelim_values, start=4):
                ws.cell(row=row_idx, column=offset, value=value if value is not None else "")
        else:
            for offset in range(4, 8):
                ws.cell(row=row_idx, column=offset, value="")
        if show_final:
            for offset, value in enumerate(final_values, start=8):
                ws.cell(row=row_idx, column=offset, value=value if value is not None else "")
        else:
            for offset in range(8, 12):
                ws.cell(row=row_idx, column=offset, value="")

        ws.cell(row=row_idx, column=12, value="")
        ws.cell(row=row_idx, column=13, value=result_item.get("rank") if result_item.get("rank") is not None else "")

    if not using_template:
        thin = Side(style="thin", color="000000")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        max_rows = max(81, len(athletes) + 2)
        for r in range(1, max_rows + 1):
            for c in range(1, 14):
                cell = ws.cell(row=r, column=c)
                cell.border = border
                cell.alignment = align
        ws.row_dimensions[1].height = 24
        ws.row_dimensions[2].height = 24
        ws.column_dimensions["A"].width = 6.5
        ws.column_dimensions["B"].width = 13
        ws.column_dimensions["C"].width = 13
        for col in ("D", "E", "F", "G", "H", "I", "J", "K", "L", "M"):
            ws.column_dimensions[col].width = 9.1
        ws.freeze_panes = "A3"
        ws.print_area = f"A1:M{max_rows}"

    max_rows = max(ws.max_row, len(athletes) + 2, 81 if not using_template else 0)
    _apply_uniform_template_font(ws, 1, max_rows, 1, 13)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def _to_chinese_group_number(value: int) -> str:
    digits = ["零", "一", "二", "三", "四", "五", "六", "七", "八", "九"]
    if value <= 0:
        return ""
    if value < 10:
        return digits[value]
    if value == 10:
        return "十"
    if value < 20:
        return "十" + digits[value % 10]
    if value < 100:
        tens = value // 10
        ones = value % 10
        return digits[tens] + "十" + (digits[ones] if ones else "")
    return str(value)


def _display_group_label(group_label: str) -> str:
    text = str(group_label or "").strip().replace("组", "")
    if not text:
        return ""
    if len(text) == 1 and "A" <= text.upper() <= "Z":
        return _to_chinese_group_number(ord(text.upper()) - ord("A") + 1)
    if text.isdigit():
        return _to_chinese_group_number(int(text))
    return text


def _group_label_sort_value(group_label: str) -> int:
    text = _display_group_label(group_label)
    if not text:
        return 999
    digit_map = {"零": 0, "一": 1, "二": 2, "三": 3, "四": 4, "五": 5, "六": 6, "七": 7, "八": 8, "九": 9}
    if text in digit_map:
        return digit_map[text]
    if text == "十":
        return 10
    if "十" in text:
        left, right = text.split("十", 1)
        tens = digit_map.get(left, 1 if left == "" else 0)
        ones = digit_map.get(right, 0 if right == "" else 0)
        return tens * 10 + ones
    return 999


def _index_to_group_label(index: int) -> str:
    return f"{_to_chinese_group_number(index + 1)}\u7ec4"


def _build_track_results_template_groups(group_rows: List[Grouping], athletes: List[Athlete]) -> List[Dict]:
    if group_rows:
        grouped: Dict[str, List[Grouping]] = {}
        for row in group_rows:
            source_label = (row.group_label or "").strip() or "default"
            grouped.setdefault(source_label, []).append(row)
        return [
            {
                "label": f"{_display_group_label(source_label)}\u7ec4",
                "rows": rows,
            }
            for source_label, rows in sorted(grouped.items(), key=lambda item: (_group_label_sort_value(item[0]), str(item[0])))
        ]

    fallback_groups: List[Dict] = []
    for idx in range(0, len(athletes), 8):
        chunk = athletes[idx : idx + 8]
        fallback_groups.append(
            {
                "label": _index_to_group_label(len(fallback_groups)),
                "rows": [
                    {
                        "student_id": _row_value(athlete, "student_id", ""),
                        "name": _row_value(athlete, "name", ""),
                        "college": _row_value(athlete, "college", ""),
                        "lane": lane_idx + 1,
                    }
                    for lane_idx, athlete in enumerate(chunk)
                ],
            }
        )

    return fallback_groups or [{"label": _index_to_group_label(0), "rows": []}]


def _track_group_export_entries(
    group: Dict,
    athletes: List[Athlete],
    result_by_student: Dict[str, Dict],
) -> List[Dict]:
    athlete_map = {
        str(_row_value(athlete, "student_id", "")).strip(): athlete
        for athlete in athletes
        if str(_row_value(athlete, "student_id", "")).strip()
    }
    entries: List[Dict] = []
    for row in group.get("rows", []):
        student_key = str(_row_value(row, "student_id", "")).strip()
        if not student_key:
            continue
        result_item = result_by_student.get(student_key)
        if not result_item:
            continue
        source = athlete_map.get(student_key) or row
        score_value = result_item.get("score")
        if not is_rankable_score(score_value):
            continue
        entries.append(
            {
                "student_id": student_key,
                "name": _row_value(source, "name", ""),
                "college": _row_value(source, "college", ""),
                "score": float(score_value),
                "rank": result_item.get("rank"),
                "points": _points_for_export_result(result_item),
            }
        )

    entries.sort(
        key=lambda item: (
            item.get("rank") if item.get("rank") is not None else 9999,
            item.get("score") if item.get("score") is not None else float("inf"),
            str(item.get("name") or ""),
            str(item.get("student_id") or ""),
        )
    )
    return entries


LONG_DISTANCE_TRACK_FOOTER_TEXT = "总裁判：             裁判长：                裁判员：               记录员："


def _numeric_group_label_text(group_label: str, fallback_index: int) -> str:
    order = _group_label_sort_value(group_label)
    if order != 999:
        return str(order)
    raw_text = str(group_label or "").strip().replace("组", "")
    return raw_text or str(fallback_index + 1)


def _build_long_distance_track_results_template_stream(
    event: Optional[str],
    group_rows: List[Grouping],
    athletes: List[Athlete],
    result_by_student: Optional[Dict[str, Dict]] = None,
    round_value: Optional[str] = None,
    gender: Optional[str] = None,
) -> io.BytesIO:
    result_by_student = result_by_student or {}
    groups = _build_track_results_template_groups(group_rows, athletes)
    has_result_export = bool(result_by_student)
    wb = Workbook()
    ws = wb.active
    ws.title = "long_distance_template"

    ws.column_dimensions["A"].width = 10
    for col_idx in range(2, 10):
        ws.column_dimensions[get_column_letter(col_idx)].width = 14

    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    block_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    title_align = Alignment(horizontal="center", vertical="center", wrap_text=False, shrink_to_fit=True)

    current_row = 1
    event_text = (event or "长跑项目").strip() or "长跑项目"
    round_text = _round_display_text(round_value)
    gender_text = _gender_title_label(gender)
    title_prefix = f"{gender_text}{event_text}{round_text}计时表" if round_text else f"{gender_text}{event_text}计时表"

    for group_index, group in enumerate(groups):
        group_text = str(group.get("label") or _index_to_group_label(group_index)).strip() or _index_to_group_label(group_index)
        group_number_text = _numeric_group_label_text(group_text, group_index)
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=9)
        title_cell = ws.cell(current_row, 1)
        if has_result_export:
            title_cell.value = f"第{group_number_text}组    {event_text}    成绩单"
        else:
            title_cell.value = f"{title_prefix}                                   第{group_number_text}组"
        title_cell.alignment = title_align
        ws.row_dimensions[current_row].height = 28

        if has_result_export:
            ordered_rows = _track_group_export_entries(group, athletes, result_by_student)
        else:
            ordered_rows = sorted(
                group.get("rows", []),
                key=lambda row: (
                    _row_value(row, "lane", 9999) if isinstance(_row_value(row, "lane", None), int) else 9999,
                    str(_row_value(row, "name", "")),
                    str(_row_value(row, "student_id", "")),
                ),
            )
        chunks = [ordered_rows[i : i + 8] for i in range(0, len(ordered_rows), 8)] or [[]]

        for chunk_index, chunk in enumerate(chunks):
            block_start = current_row + 1
            labels = ["名次", "姓名", "单位", "成绩", "备注"] if has_result_export else ["号码", "姓名", "单位", "成绩", "备注"]
            values_by_row = [[] for _ in labels]

            for item in chunk:
                if has_result_export:
                    values_by_row[0].append(_row_value(item, "rank", ""))
                    values_by_row[1].append(_row_value(item, "name", ""))
                    values_by_row[2].append(_row_value(item, "college", ""))
                    values_by_row[3].append(_row_value(item, "score", ""))
                    values_by_row[4].append(_row_value(item, "points", 0))
                else:
                    student_key = str(_row_value(item, "student_id", "")).strip()
                    result_item = result_by_student.get(student_key, {})
                    lane_value = _row_value(item, "lane", "")
                    values_by_row[0].append("" if lane_value is None else lane_value)
                    values_by_row[1].append(_row_value(item, "name", ""))
                    values_by_row[2].append(_row_value(item, "college", ""))
                    values_by_row[3].append(result_item.get("score") if result_item.get("score") is not None else "")
                    values_by_row[4].append("")

            for row_offset, label in enumerate(labels):
                row_number = block_start + row_offset
                ws.row_dimensions[row_number].height = 24
                label_cell = ws.cell(row_number, 1, label)
                label_cell.border = border
                label_cell.alignment = block_align
                row_values = values_by_row[row_offset]
                for col_offset in range(8):
                    column_number = 2 + col_offset
                    value = row_values[col_offset] if col_offset < len(row_values) else ""
                    data_cell = ws.cell(row_number, column_number, value)
                    data_cell.border = border
                    data_cell.alignment = block_align

            current_row = block_start + 4
            if chunk_index < len(chunks) - 1:
                current_row += 1
                ws.row_dimensions[current_row].height = 10

        current_row += 1
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=9)
        footer_cell = ws.cell(current_row, 1, LONG_DISTANCE_TRACK_FOOTER_TEXT)
        footer_cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=False)
        ws.row_dimensions[current_row].height = 24

        if group_index < len(groups) - 1:
            current_row += 1
            ws.row_dimensions[current_row].height = 12
            current_row += 1

    _apply_uniform_template_font(ws, 1, ws.max_row, 1, 9)
    ws.print_area = f"A1:I{ws.max_row}"

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def _build_track_results_template_stream(
    event: Optional[str],
    group_rows: List[Grouping],
    athletes: List[Athlete],
    result_by_student: Optional[Dict[str, Dict]] = None,
    round_value: Optional[str] = None,
    gender: Optional[str] = None,
) -> io.BytesIO:
    if is_long_distance_track_event(event):
        return _build_long_distance_track_results_template_stream(
            event,
            group_rows,
            athletes,
            result_by_student=result_by_student,
            round_value=round_value,
            gender=gender,
        )

    result_by_student = result_by_student or {}
    groups = _build_track_results_template_groups(group_rows, athletes)
    has_result_export = bool(result_by_student)
    template_path = os.path.join(
        os.path.dirname(__file__),
        "templates",
        "results_track_style_template.xlsx",
    )
    if os.path.exists(template_path):
        template_wb = load_workbook(template_path)
        template_ws = template_wb.active
        wb = Workbook()
        ws = wb.active
        ws.title = template_ws.title
        ws.sheet_properties = copy(template_ws.sheet_properties)
        ws.sheet_format = copy(template_ws.sheet_format)
        ws.page_margins = copy(template_ws.page_margins)
        ws.page_setup = copy(template_ws.page_setup)
        ws.print_options = copy(template_ws.print_options)
        ws.freeze_panes = template_ws.freeze_panes

        for col_idx in range(1, 10):
            letter = get_column_letter(col_idx)
            ws.column_dimensions[letter].width = template_ws.column_dimensions[letter].width

        footer_text = template_ws.cell(7, 1).value or ""
        footer_style = template_ws.cell(7, 1)
        title_style = template_ws.cell(1, 1)
        lane_headers = (
            [
                "\u540d\u6b21",
                "\u7b2c\u4e00\u540d",
                "\u7b2c\u4e8c\u540d",
                "\u7b2c\u4e09\u540d",
                "\u7b2c\u56db\u540d",
                "\u7b2c\u4e94\u540d",
                "\u7b2c\u516d\u540d",
                "\u7b2c\u4e03\u540d",
                "\u7b2c\u516b\u540d",
            ]
            if has_result_export
            else [
                "\u9053\u6b21",
                "\u7b2c\u4e00\u9053",
                "\u7b2c\u4e8c\u9053",
                "\u7b2c\u4e09\u9053",
                "\u7b2c\u56db\u9053",
                "\u7b2c\u4e94\u9053",
                "\u7b2c\u516d\u9053",
                "\u7b2c\u4e03\u9053",
                "\u7b2c\u516b\u9053",
            ]
        )
        gender_text = _gender_title_label(gender)

        for block_idx, group in enumerate(groups):
            start_row = 1 + block_idx * 8

            if template_ws.row_dimensions[1].height is not None:
                ws.row_dimensions[start_row].height = template_ws.row_dimensions[1].height
            for row_offset in range(1, 6):
                src_row = 1 + row_offset
                dst_row = start_row + row_offset
                if template_ws.row_dimensions[src_row].height is not None:
                    ws.row_dimensions[dst_row].height = template_ws.row_dimensions[src_row].height
                for col_idx in range(1, 10):
                    src_cell = template_ws.cell(src_row, col_idx)
                    dst_cell = ws.cell(dst_row, col_idx)
                    dst_cell._style = copy(src_cell._style)
                    dst_cell.font = copy(src_cell.font)
                    dst_cell.fill = copy(src_cell.fill)
                    dst_cell.border = copy(src_cell.border)
                    dst_cell.alignment = copy(src_cell.alignment)
                    dst_cell.number_format = src_cell.number_format
                    dst_cell.protection = copy(src_cell.protection)
                    dst_cell.value = src_cell.value

            if template_ws.row_dimensions[7].height is not None:
                ws.row_dimensions[start_row + 6].height = template_ws.row_dimensions[7].height

            ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=9)
            ws.merge_cells(start_row=start_row + 6, start_column=1, end_row=start_row + 6, end_column=9)

            title_cell = ws.cell(start_row, 1)
            title_cell._style = copy(title_style._style)
            title_cell.font = copy(title_style.font)
            title_cell.fill = copy(title_style.fill)
            title_cell.border = copy(title_style.border)
            title_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False, shrink_to_fit=True)
            title_cell.number_format = title_style.number_format
            title_cell.protection = copy(title_style.protection)

            group_text = str(group.get("label") or _index_to_group_label(block_idx)).strip() or _index_to_group_label(block_idx)
            event_text = (event or "\u5f84\u8d5b\u9879\u76ee").strip() or "\u5f84\u8d5b\u9879\u76ee"
            round_text = _round_display_text(round_value)
            title_prefix = f"{event_text}{round_text}\u8ba1\u65f6\u8868" if round_text else f"{event_text}\u8ba1\u65f6\u8868"
            group_display = _display_group_label(group_text)
            ordinal_group_text = f"\u7b2c{group_display}\u7ec4" if group_display else group_text
            title_parts = (
                [part for part in [ordinal_group_text, event_text, "\u6210\u7ee9\u5355"] if part]
                if has_result_export
                else [part for part in [gender_text, title_prefix, ordinal_group_text] if part]
            )
            title_cell.value = "    ".join(title_parts)

            for col_idx, header_value in enumerate(lane_headers, start=1):
                ws.cell(start_row + 1, col_idx).value = header_value

            footer_cell = ws.cell(start_row + 6, 1)
            footer_cell._style = copy(footer_style._style)
            footer_cell.font = copy(footer_style.font)
            footer_cell.fill = copy(footer_style.fill)
            footer_cell.border = copy(footer_style.border)
            footer_cell.alignment = copy(footer_style.alignment)
            footer_cell.number_format = footer_style.number_format
            footer_cell.protection = copy(footer_style.protection)
            footer_cell.value = footer_text

            if has_result_export:
                ranked_rows = _track_group_export_entries(group, athletes, result_by_student)
                for index, item in enumerate(ranked_rows[:8], start=2):
                    ws.cell(start_row + 2, index).value = item.get("name", "")
                    ws.cell(start_row + 3, index).value = item.get("college", "")
                    ws.cell(start_row + 4, index).value = item.get("score", "")
                    ws.cell(start_row + 5, index).value = item.get("points", 0)
            else:
                used_columns = set()
                next_col = 2
                for row in group.get("rows", []):
                    lane = _row_value(row, "lane", None)
                    name = _row_value(row, "name", "")
                    college = _row_value(row, "college", "")
                    student_key = str(_row_value(row, "student_id", "")).strip()
                    result_item = result_by_student.get(student_key, {})

                    column = None
                    if isinstance(lane, int) and 1 <= lane <= 8:
                        candidate = lane + 1
                        if candidate not in used_columns:
                            column = candidate

                    if column is None:
                        while next_col in used_columns and next_col <= 9:
                            next_col += 1
                        if next_col > 9:
                            continue
                        column = next_col

                    used_columns.add(column)
                    ws.cell(start_row + 2, column).value = name
                    ws.cell(start_row + 3, column).value = college
                    ws.cell(start_row + 4, column).value = result_item.get("score") if result_item.get("score") is not None else ""
                    ws.cell(start_row + 5, column).value = ""

            if block_idx < len(groups) - 1 and template_ws.row_dimensions[8].height is not None:
                ws.row_dimensions[start_row + 7].height = template_ws.row_dimensions[8].height
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "track_template"
        headers = (
            [
                "\u540d\u6b21",
                "\u7b2c\u4e00\u540d",
                "\u7b2c\u4e8c\u540d",
                "\u7b2c\u4e09\u540d",
                "\u7b2c\u56db\u540d",
                "\u7b2c\u4e94\u540d",
                "\u7b2c\u516d\u540d",
                "\u7b2c\u4e03\u540d",
                "\u7b2c\u516b\u540d",
            ]
            if has_result_export
            else [
                "\u9053\u6b21",
                "\u7b2c\u4e00\u9053",
                "\u7b2c\u4e8c\u9053",
                "\u7b2c\u4e09\u9053",
                "\u7b2c\u56db\u9053",
                "\u7b2c\u4e94\u9053",
                "\u7b2c\u516d\u9053",
                "\u7b2c\u4e03\u9053",
                "\u7b2c\u516b\u9053",
            ]
        )
        labels = ["\u59d3\u540d", "\u5355\u4f4d", "\u6210\u7ee9", "\u5907\u6ce8"]
        thin = Side(style="thin", color="000000")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        align = Alignment(horizontal="center", vertical="center", wrap_text=True)

        ws.merge_cells("A1:I1")
        round_text = _round_display_text(round_value)
        default_event_text = "\u5f84\u8d5b\u9879\u76ee"
        title_prefix = f"{(event or default_event_text)}{round_text}\u8ba1\u65f6\u8868" if round_text else (event or default_event_text) + "\u8ba1\u65f6\u8868"
        gender_text = _gender_title_label(gender)
        title_parts = (
            [part for part in ["\u7b2c\u4e00\u7ec4", event or default_event_text, "\u6210\u7ee9\u5355"] if part]
            if has_result_export
            else [part for part in [gender_text, title_prefix, "\u7b2c\u4e00\u7ec4"] if part]
        )
        ws["A1"] = "    ".join(title_parts)
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=False, shrink_to_fit=True)

        for i, h in enumerate(headers, start=1):
            cell = ws.cell(row=2, column=i, value=h)
            cell.border = border
            cell.alignment = align

        for r, name in enumerate(labels, start=3):
            for c in range(1, 10):
                cell = ws.cell(row=r, column=c)
                cell.border = border
                cell.alignment = align
                if c == 1:
                    cell.value = name

    if os.path.exists(template_path):
        ws.print_area = f"A1:I{max(7, len(groups) * 8 - 1)}"

    max_rows = max(ws.max_row, len(groups) * 8 - 1 if groups else ws.max_row)
    _apply_uniform_template_font(ws, 1, max_rows, 1, 9)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


@app.get("/export/results-template")
def export_results_template(
    event: Optional[str] = None,
    gender: Optional[str] = None,
    round: Optional[str] = None,
    db: Session = Depends(get_db),
):
    if not event:
        raise HTTPException(status_code=400, detail="请先选择项目后再导出模板")

    athletes = _list_template_athletes(event, db, gender=gender)
    event_kind = classify_event_table(event) if event else None
    if event_kind == "jump":
        output = _build_jump_results_template_stream(athletes)
        timestamp = get_timestamp()
        filename = build_results_template_filename(event, timestamp, gender)
        return _build_excel_stream_response(output, filename)

    if event_kind == "distance":
        output = _build_field_results_template_stream(athletes)
        timestamp = get_timestamp()
        filename = build_results_template_filename(event, timestamp, gender)
        return _build_excel_stream_response(output, filename)

    if event_kind == "track":
        group_rows, template_round = _list_template_groupings(event, round, db, gender=gender)
        output = _build_track_results_template_stream(event, group_rows, athletes, round_value=template_round or round, gender=gender)
        timestamp = get_timestamp()
        filename = build_results_template_filename(event, timestamp, gender)
        return _build_excel_stream_response(output, filename)

    raise HTTPException(status_code=400, detail="当前项目暂不支持该模板导出")


def _round_priority_for_ranking(round_value: str | None) -> int:
    value = (round_value or "").lower()
    if value == "final":
        return 2
    if value == "one":
        return 1
    return 0


def _should_replace_round_row(prev_row, new_row) -> bool:
    prev_priority = _round_priority_for_ranking(getattr(prev_row, "round", None))
    new_priority = _round_priority_for_ranking(getattr(new_row, "round", None))
    if new_priority != prev_priority:
        return new_priority > prev_priority
    prev_date = getattr(prev_row, "date", None) or datetime.min
    new_date = getattr(new_row, "date", None) or datetime.min
    return new_date >= prev_date


def _gender_bucket(value: Optional[str]) -> Optional[str]:
    raw = (value or "").strip()
    text = raw.lower()
    if not text:
        return None
    if "\u7537" in raw or text in {"male", "m", "man", "men"}:
        return "male"
    if "\u5973" in raw or text in {"female", "f", "woman", "women"}:
        return "female"
    return None


def _team_manual_bonus_map(db: Session) -> Dict:
    return {
        (row.student_id, row.event): row
        for row in db.scalars(select(TeamRankingManualBonus)).all()
    }


def _collect_team_ranking_records(db: Session) -> Dict[str, List[Dict]]:
    athlete_map = {
        (row.student_id, row.event): row
        for row in db.scalars(select(Athlete)).all()
    }
    bonus_map = _team_manual_bonus_map(db)

    relay_rows = db.scalars(
        select(Result).where(Result.round.in_(["final", "one"]))
    ).all()
    relay_map: Dict = {}
    for row in relay_rows:
        if not is_relay_event(row.event):
            continue
        athlete = athlete_map.get((row.student_id, row.event))
        gender_bucket = _gender_bucket(getattr(athlete, "gender", None))
        if gender_bucket not in {"male", "female"}:
            continue
        key = (gender_bucket, row.event, row.student_id)
        prev = relay_map.get(key)
        if prev is None or _should_replace_round_row(prev, row):
            relay_map[key] = row

    by_event: Dict = {}
    for (gender_bucket, event_name, student_id), row in relay_map.items():
        if row.score is None:
            continue
        by_event.setdefault((gender_bucket, event_name), []).append(
            {"student_id": student_id, "score": row.score}
        )

    rank_by_event: Dict = {}
    for key, rows in by_event.items():
        ranked = rank_athletes_by_score(rows, key[1])
        rank_by_event[key] = {
            entry["student_id"]: entry["rank"] for entry in ranked
        }

    grouped = {"male": [], "female": []}
    for (gender_bucket, event_name, student_id), row in relay_map.items():
        if row.score is None:
            continue
        rank_value = rank_by_event.get((gender_bucket, event_name), {}).get(student_id)
        if rank_value is None:
            rank_value = int(row.rank or 0)
        if rank_value <= 0:
            continue
        athlete = athlete_map.get((student_id, event_name))
        record_broken = bool(getattr(athlete, "record_broken", False))
        base_points = points_for_relay_team_rank(int(rank_value)) + points_for_record_broken(record_broken)
        manual_bonus = int(getattr(bonus_map.get((student_id, event_name)), "manual_points", 0) or 0)
        grouped[gender_bucket].append(
            {
                "student_id": student_id,
                "event": event_name,
                "college": row.college,
                "score": float(row.score),
                "rank": int(rank_value),
                "base_points": base_points,
                "manual_bonus": manual_bonus,
                "points": base_points + manual_bonus,
                "total_points": base_points + manual_bonus,
            }
        )

    for bucket in grouped.values():
        bucket.sort(key=lambda item: (item["event"], item["rank"], item["college"], item["student_id"]))
    return grouped


def _collect_ranking_records(db: Session) -> List[Dict]:
    records: List[Dict] = []
    athlete_record_flags = {
        (row.student_id, row.event): bool(row.record_broken)
        for row in db.scalars(select(Athlete)).all()
    }

    track_rows = db.scalars(
        select(Result).where(Result.round.in_(["final", "one"]))
    ).all()
    track_map: Dict = {}
    for row in track_rows:
        if classify_event_table(row.event) != "track":
            continue
        if is_relay_event(row.event):
            continue
        key = (row.student_id, row.event)
        prev = track_map.get(key)
        if prev is None or _should_replace_round_row(prev, row):
            track_map[key] = row
    for row in track_map.values():
        records.append(
            {
                "student_id": row.student_id,
                "name": row.name,
                "college": row.college,
                "event": row.event,
                "score": row.score,
                "rank": row.rank,
                "record_broken": athlete_record_flags.get((row.student_id, row.event), False),
            }
        )

    jump_rows = db.scalars(
        select(ResultJump).where(ResultJump.round.in_(["final", "one"]))
    ).all()
    jump_map: Dict = {}
    for row in jump_rows:
        if classify_event_table(row.event) != "jump":
            continue
        key = (row.student_id, row.event)
        prev = jump_map.get(key)
        if prev is None or _should_replace_round_row(prev, row):
            jump_map[key] = row
    for row in jump_map.values():
        records.append(
            {
                "student_id": row.student_id,
                "name": row.name,
                "college": row.college,
                "event": row.event,
                "score": row.score,
                "rank": row.rank,
                "record_broken": athlete_record_flags.get((row.student_id, row.event), False),
            }
        )

    distance_rows = db.scalars(
        select(ResultDistance).where(ResultDistance.final_best.isnot(None))
    ).all()
    for row in distance_rows:
        if classify_event_table(row.event) != "distance":
            continue
        records.append(
            {
                "student_id": row.student_id,
                "name": row.name,
                "college": row.college,
                "event": row.event,
                "score": row.final_best,
                "rank": row.rank,
                "record_broken": athlete_record_flags.get((row.student_id, row.event), False),
            }
        )

    by_event: Dict[str, List[Dict]] = {}
    for item in records:
        score = item.get("score")
        if not is_rankable_score(score):
            continue
        by_event.setdefault(item["event"], []).append(
            {"student_id": item["student_id"], "score": score}
        )

    rank_by_event: Dict[str, Dict[str, int]] = {}
    for event_name, rows in by_event.items():
        ranked = rank_athletes_by_score(rows, event_name)
        rank_by_event[event_name] = {
            entry["student_id"]: entry["rank"] for entry in ranked
        }

    normalized: List[Dict] = []
    for item in records:
        score = item.get("score")
        if not is_rankable_score(score):
            continue
        rank_value = rank_by_event.get(item["event"], {}).get(item["student_id"])
        if rank_value is None:
            rank_value = int(item.get("rank") or 0)
        if rank_value <= 0:
            continue
        normalized.append(
            {
                "student_id": item["student_id"],
                "name": item["name"],
                "college": item["college"],
                "event": item["event"],
                "score": float(score),
                "rank": int(rank_value),
                "points": points_for_rank(int(rank_value)) + points_for_record_broken(bool(item.get("record_broken"))),
            }
        )

    return normalized


def rebuild_personal_rankings(db: Session) -> int:
    records = _collect_ranking_records(db)
    db.execute(delete(PersonalRanking))

    total_by_student: Dict[str, int] = {}
    for item in records:
        sid = item["student_id"]
        total_by_student[sid] = total_by_student.get(sid, 0) + int(item["points"])

    for item in records:
        db.add(
            PersonalRanking(
                student_id=item["student_id"],
                name=item["name"],
                college=item["college"],
                event=item["event"],
                score=item["score"],
                rank=item["rank"],
                points=item["points"],
                total_points=total_by_student.get(item["student_id"], 0),
            )
        )

    db.commit()
    return len(records)


def rebuild_college_rankings(db: Session) -> int:
    rows = db.scalars(select(PersonalRanking)).all()
    db.execute(delete(CollegeRanking))

    college_stats: Dict[str, Dict] = {}
    for row in rows:
        bucket = college_stats.setdefault(
            row.college,
            {"events": set(), "students": set(), "total_points": 0},
        )
        bucket["events"].add(row.event)
        bucket["students"].add(row.student_id)
        bucket["total_points"] += int(row.points)

    for college, stats in college_stats.items():
        db.add(
            CollegeRanking(
                college=college,
                event_count=len(stats["events"]),
                participant_count=len(stats["students"]),
                total_points=int(stats["total_points"]),
            )
        )

    db.commit()
    return len(college_stats)


@app.post("/rankings/personal/initialize")
def initialize_personal_rankings(db: Session = Depends(get_db)):
    personal_count = rebuild_personal_rankings(db)
    return {"message": "initialized", "count": personal_count}


@app.post("/rankings/college/initialize")
def initialize_college_rankings(db: Session = Depends(get_db)):
    if db.scalar(select(PersonalRanking.id).limit(1)) is None:
        rebuild_personal_rankings(db)
    college_count = rebuild_college_rankings(db)
    return {"message": "initialized", "count": college_count}


@app.get("/rankings/personal/list")
def list_personal_rankings(db: Session = Depends(get_db)):
    query = select(PersonalRanking).order_by(
        PersonalRanking.total_points.desc(),
        PersonalRanking.points.desc(),
        PersonalRanking.rank.asc(),
        PersonalRanking.student_id.asc(),
        PersonalRanking.event.asc(),
    )
    rows = [
        row
        for row in db.scalars(query).all()
        if row.rank is not None and int(row.rank) <= 8
    ]
    return [to_dict(row) for row in rows]


@app.get("/rankings/college/list")
def list_college_rankings(db: Session = Depends(get_db)):
    query = select(CollegeRanking).order_by(
        CollegeRanking.total_points.desc(),
        CollegeRanking.event_count.desc(),
        CollegeRanking.participant_count.desc(),
        CollegeRanking.college.asc(),
    )
    rows = db.scalars(query).all()
    return [to_dict(row) for row in rows]

@app.get("/rankings/team/list")
def list_team_rankings(db: Session = Depends(get_db)):
    return _collect_team_ranking_records(db)


@app.post("/rankings/team/bonus")
def save_team_ranking_bonus(req: TeamRankingManualBonusRequest, db: Session = Depends(get_db)):
    athlete = db.scalar(
        select(Athlete).where(
            Athlete.student_id == req.student_id,
            Athlete.event == req.event,
        )
    )
    if not athlete or not is_relay_event(req.event):
        raise HTTPException(status_code=404, detail="Team ranking row not found")

    existing = db.scalar(
        select(TeamRankingManualBonus).where(
            TeamRankingManualBonus.student_id == req.student_id,
            TeamRankingManualBonus.event == req.event,
        )
    )

    manual_points = int(req.manual_points or 0)
    if manual_points == 0:
        if existing:
            db.delete(existing)
            db.commit()
        return {"message": "saved", "manual_points": 0}

    if not existing:
        existing = TeamRankingManualBonus(
            student_id=req.student_id,
            event=req.event,
            manual_points=manual_points,
        )
        db.add(existing)
    else:
        existing.manual_points = manual_points

    db.commit()
    return {"message": "saved", "manual_points": manual_points}


@app.get("/export/personal-rankings")
def export_personal_rankings(db: Session = Depends(get_db)):
    rows = [
        row
        for row in db.scalars(
            select(PersonalRanking).order_by(
                PersonalRanking.event.asc(),
                PersonalRanking.rank.asc(),
                PersonalRanking.student_id.asc(),
            )
        ).all()
        if row.rank is not None and int(row.rank) <= 8
    ]
    columns = ["学号", "姓名", "学院", "项目", "成绩", "排名", "积分", "总积分"]
    data = [
        {
            "\u5b66\u53f7": r.student_id,
            "\u59d3\u540d": r.name,
            "\u5b66\u9662": r.college,
            "\u9879\u76ee": r.event,
            "\u6210\u7ee9": r.score,
            "\u6392\u540d": r.rank,
            "\u79ef\u5206": r.points,
            "\u603b\u79ef\u5206": r.total_points,
        }
        for r in rows
    ]
    timestamp = get_timestamp()
    return excel_response(data, f"personal_rankings_{timestamp}.xlsx", columns=columns)











@app.get("/export/college-rankings")
def export_college_rankings(db: Session = Depends(get_db)):
    rows = db.scalars(select(CollegeRanking)).all()
    columns = ["学院", "项目数", "参赛人数", "总积分"]
    data = [
        {
            "\u5b66\u9662": r.college,
            "\u9879\u76ee\u6570": r.event_count,
            "\u53c2\u8d5b\u4eba\u6570": r.participant_count,
            "\u603b\u79ef\u5206": r.total_points,
        }
        for r in rows
    ]
    timestamp = get_timestamp()
    return excel_response(data, f"college_rankings_{timestamp}.xlsx", columns=columns)











@app.get("/export/team-rankings")
def export_team_rankings(db: Session = Depends(get_db)):
    grouped = _collect_team_ranking_records(db)
    columns = ["团体类别", "项目", "学院", "成绩", "排名", "项目得分", "额外加分", "总积分"]
    data: List[Dict] = []

    for bucket_key, bucket_label in (("male", "男子团体"), ("female", "女子团体")):
        for row in grouped.get(bucket_key, []):
            data.append(
                {
                    "团体类别": bucket_label,
                    "项目": row.get("event"),
                    "学院": row.get("college"),
                    "成绩": row.get("score"),
                    "排名": row.get("rank"),
                    "项目得分": row.get("base_points", row.get("points", 0)),
                    "额外加分": row.get("manual_bonus", 0),
                    "总积分": row.get("total_points", row.get("points", 0)),
                }
            )

    timestamp = get_timestamp()
    return excel_response(data, f"team_rankings_{timestamp}.xlsx", columns=columns)


def format_group_name(gender: str, group_label: str) -> str:
    value = (gender or "").strip().lower()
    display_label = _display_group_label(group_label)
    male_labels = {"\u7537", "male", "\u7537\u5b50"}
    female_labels = {"\u5973", "female", "\u5973\u5b50"}
    if value in male_labels:
        return f"\u7537\u5b50{display_label}\u7ec4"
    if value in female_labels:
        return f"\u5973\u5b50{display_label}\u7ec4"
    return f"\u6df7\u5408{display_label}\u7ec4"

def excel_response(data: List[Dict], filename: str, columns: Optional[List[str]] = None):
    df = pd.DataFrame(data, columns=columns) if columns else pd.DataFrame(data)
    output = io.BytesIO()
    df.to_excel(output, index=False)
    output.seek(0)
    safe_name = filename.encode("ascii", "ignore").decode("ascii") or "export.xlsx"
    encoded_name = urllib.parse.quote(filename)
    headers = {
        "Content-Disposition": f"attachment; filename={safe_name}; filename*=UTF-8''{encoded_name}"
    }
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers,
    )













def update_final_results_distance(event_name: str, db: Session, at_time: datetime | None = None):
    rows = db.scalars(
        select(ResultDistance).where(ResultDistance.event == event_name)
    ).all()
    scored = []
    for row in rows:
        if row.final_best is None:
            continue
        scored.append({"student_id": row.student_id, "score": row.final_best})
    ranked = rank_athletes_by_score(scored, event_name)
    db.execute(delete(Result).where(Result.event == event_name))
    best_candidate = ranked[0] if ranked else None
    event_row = db.scalar(select(Event).where(Event.name == event_name))
    if best_candidate:
        if not event_row:
            event_row = Event(name=event_name, participant_count=0, best_record=None, record_holder=None)
            db.add(event_row)
        should_update = event_row.best_record is None or best_candidate["score"] > event_row.best_record
        if should_update:
            athlete_best = db.scalar(
                select(Athlete).where(
                    Athlete.student_id == best_candidate["student_id"],
                    Athlete.event == event_name,
                )
            )
            event_row.best_record = best_candidate["score"]
            event_row.record_holder = athlete_best.name if athlete_best else ""
            if athlete_best:
                athlete_best.record_broken = True
    now = at_time or datetime.utcnow()
    for row in ranked:
        athlete = db.scalar(
            select(Athlete).where(
                Athlete.student_id == row["student_id"],
                Athlete.event == event_name,
            )
        )
        if not athlete:
            continue
        result = Result(
            event=event_name,
            name=athlete.name,
            student_id=athlete.student_id,
            college=athlete.college,
            score=row["score"],
            rank=row["rank"],
            date=now,
            round="final",
        )
        db.add(result)
        athlete.score = row["score"]
        athlete.rank = row["rank"]

def update_final_results(event_name: str, db: Session):



    group_rows = db.scalars(



        select(Grouping).where(



            Grouping.event == event_name,



        )



    ).all()



    final_scores = []



    for row in group_rows:



        score = row.final_score if row.final_score is not None else (



            row.prelim_score if row.round == "one" else None



        )



        if score is not None:



            final_scores.append(



                {



                    "student_id": row.student_id,



                    "score": score,



                }



            )



    ranked = rank_athletes_by_score(final_scores, event_name)



    db.execute(delete(Result).where(Result.event == event_name))



    best_candidate = ranked[0] if ranked else None



    event_row = db.scalar(select(Event).where(Event.name == event_name))



    if best_candidate:



        if not event_row:



            event_row = Event(



                name=event_name,



                participant_count=0,



                best_record=None,



                record_holder=None,



            )



            db.add(event_row)



        should_update = False



        if event_row.best_record is None:



            should_update = True



        else:



            if is_distance_event(event_name):



                should_update = best_candidate["score"] > event_row.best_record



            else:



                should_update = best_candidate["score"] < event_row.best_record



        if should_update:



            athlete = db.scalar(



                select(Athlete).where(



                    Athlete.student_id == best_candidate["student_id"],



                    Athlete.event == event_name,



                )



            )



            event_row.best_record = best_candidate["score"]



            event_row.record_holder = athlete.name if athlete else "鑼呰伋娼炴皳鑱磋伌姘撹仚楣挎皳濞勮姦"



    for row in ranked:



        athlete = db.scalar(



            select(Athlete).where(



                Athlete.student_id == row["student_id"],



                Athlete.event == event_name,



            )



        )



        if not athlete:



            continue



        result = Result(



            event=event_name,



            name=athlete.name,



            student_id=athlete.student_id,



            college=athlete.college,



            score=row["score"],



            rank=row["rank"],



            date=datetime.utcnow(),



            round="final",



        )



        db.add(result)



        athlete.score = row["score"]



        athlete.rank = row["rank"]



# FORCE_NORMALIZED_EVENT_SETS
MULTI_ROUND_EVENTS = {
    "\u0031\u0030\u0030\u7c73",
    "\u0032\u0030\u0030\u7c73",
    "\u0034\u0030\u0030\u7c73",
    "\u0038\u0030\u0030\u7c73",
    "\u0031\u0035\u0030\u0030\u7c73",
    "\u0033\u0030\u0030\u0030\u7c73",
    "\u0035\u0030\u0030\u0030\u7c73",
    "\u0034*\u0031\u0030\u0030\u7c73",
    "\u0034*\u0032\u0030\u0030\u7c73",
    "\u0034*\u0034\u0030\u0030\u7c73",
}
MIXED_EVENTS = {"\u6df7\u5408\u0034*\u0031\u0030\u0030\u7c73", "\u6df7\u5408\u0034*\u0034\u0030\u0030\u7c73"}
DISTANCE_KEYWORDS = {
    "\u8df3",
    "\u6295",
    "\u63b7",
    "\u94c5\u7403",
    "\u6807\u67aa",
    "\u94c1\u997c",
    "\u8df3\u8fdc",
    "\u4e09\u7ea7\u8df3\u8fdc",
    "\u7acb\u5b9a\u8df3\u8fdc",
    "\u6025\u884c\u8df3\u8fdc",
    "\u8df3\u9ad8",
}
DISTANCE_EVENTS_ALIAS = {
    "\u8df3\u8fdc",
    "\u4e09\u7ea7\u8df3\u8fdc",
    "\u7acb\u5b9a\u8df3\u8fdc",
    "\u6025\u884c\u8df3\u8fdc",
    "\u94c5\u7403",
    "\u6807\u67aa",
    "\u94c1\u997c",
}




























